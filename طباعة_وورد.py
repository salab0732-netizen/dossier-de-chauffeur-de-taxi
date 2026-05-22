#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
وحدة ملء نماذج Word وطباعتها
تعمل مع ملف: تسجيل_و_تسديد1.docx
"""
import os, platform, subprocess, io, re
from datetime import datetime
from dateutil.relativedelta import relativedelta
from utils import fmt_date, fmt_raqm, clean_wilaya, clean_takh, EXCEL_FILE, SHEETS, COLS_MARKAZ

# مسار نموذج Word
TEMPLATE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_FILE = os.path.join(TEMPLATE_DIR, "تسجيل_و_تسديد1.docx")
OUTPUT_DIR    = os.path.join(os.path.dirname(os.path.abspath(__file__)), "وثائق_Word")




def fix_all_dates_in_doc(doc):
    """تحويل كل DD/MM/YYYY في الوثيقة إلى YYYY/MM/DD"""
    import re as _re
    def convert_dates(text):
        return _re.sub(r'\b(\d{2})/(\d{2})/(\d{4})\b',
            lambda m: f"‏{m.group(3)}/{m.group(2)}/{m.group(1)}", text)
    def fix_para(para):
        if not para.runs: return
        full = "".join(r.text for r in para.runs)
        fixed = convert_dates(full)
        if fixed != full:
            para.runs[0].text = fixed
            for r in para.runs[1:]: r.text = ""
    for para in doc.paragraphs: fix_para(para)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs: fix_para(para)
    return doc


MONTANTS = {
    "أشخاص/بضائع": 32300.00,
    "أشخاص":        17000.00,
    "بضائع":        22500.00,
    "مواد خطرة":   16000.00,
}

def get_montant(takh):
    from utils import clean_takh as _ct
    t = _ct(str(takh or ""))
    for key, val in MONTANTS.items():
        if key in t:
            return val
    return 0.0

def clean_wilaya(w):
    """إزالة رقم الولاية من صيغة 32-البيض"""
    import re
    return re.sub(r'^\d+-', '', str(w or '')).strip()

def extract_ville(markaz):
    """
    استخراج اسم المدينة بالأولوية:
    1. بلدية المركز (حقل مستقل)              ← الأدق
    2. آخر جزء من عنوان المركز بعد الفاصلة  ← احتياط
    3. مديرية النقل لولاية (بعد إزالة الرقم) ← احتياط أخير
    """
    # 1. حقل بلدية المركز المستقل
    baladia = str(markaz.get("بلدية المركز", "") or "").strip()
    if baladia:
        return baladia
    # 2. آخر جزء من العنوان
    adresse = str(markaz.get("عنوان مركز التكوين", "") or "").strip()
    if adresse:
        parts = [p.strip() for p in adresse.replace("،", ",").split(",")]
        parts = [p for p in parts if p]
        if parts:
            return parts[-1]
    # 3. مديرية النقل احتياطاً
    return clean_wilaya(markaz.get("مديرية النقل لولاية", ""))

def clean_takh(takh):
    return (str(takh)
            .replace("نقل الأشخاص و البضائع", "أشخاص/بضائع")
            .replace("نقل الأشخاص",  "الأشخاص")
            .replace("نقل البضائع",   "البضائع")
            .replace("نقل المواد الخطرة", "المواد الخطرة")
            .replace("الأشخاص و البضائع", "أشخاص/بضائع"))

def make_qr(text):
    """توليد QR Code وإرجاعه كـ BytesIO"""
    import qrcode
    qr = qrcode.QRCode(version=2,
                       error_correction=qrcode.constants.ERROR_CORRECT_M,
                       box_size=4, border=1)
    qr.add_data(text); qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO(); img.save(buf, format="PNG"); buf.seek(0)
    return buf

def replace_in_para(para, replacements):
    """استبدال العلامات في فقرة Word — يعالج كل run على حدة للحفاظ على التنسيق"""
    # المرحلة 1: فحص سريع — هل توجد أي علامة في الفقرة كلها؟
    full = "".join(r.text or "" for r in para.runs)
    if not full: return
    if not any(key in full for key in replacements): return

    # المرحلة 2: تجميع نص الفقرة كاملاً مع تتبع حدود كل run
    # أحياناً تكون العلامة مُقسَّمة على runs متعددة بسبب Word
    # نبني قائمة (start, end, run_index) لكل run
    positions = []
    pos = 0
    for i, r in enumerate(para.runs):
        t = r.text or ""
        positions.append((pos, pos + len(t), i))
        pos += len(t)

    # المرحلة 3: استبدال في النص الكامل
    new_full = full
    # الأطول أولاً لتجنب التداخل
    for key in sorted(replacements.keys(), key=len, reverse=True):
        new_full = new_full.replace(key, str(replacements[key]))

    if new_full == full:
        return  # لم يتغير شيء

    # المرحلة 4: إعادة توزيع النص على الـ runs مع الحفاظ على التنسيق
    # نتحقق أولاً إذا كانت كل علامة داخل run واحد فقط
    # (الحالة الشائعة) — نستبدل مباشرة في كل run
    all_single = True
    for key in replacements:
        if key in full:
            # ابحث عن موضع العلامة في النص الكامل
            idx = full.find(key)
            while idx != -1:
                # هل العلامة كلها داخل run واحد؟
                in_single = False
                for (start, end, ri) in positions:
                    if start <= idx and (idx + len(key)) <= end:
                        in_single = True
                        break
                if not in_single:
                    all_single = False
                    break
                idx = full.find(key, idx + 1)
        if not all_single:
            break

    if all_single:
        # كل علامة داخل run واحد — استبدال مباشر مع الحفاظ على التنسيق
        for r in para.runs:
            if not r.text: continue
            new_text = r.text
            for key in sorted(replacements.keys(), key=len, reverse=True):
                new_text = new_text.replace(key, str(replacements[key]))
            r.text = new_text
    else:
        # علامة مُقسَّمة على runs متعددة — نضطر لدمج الكل في أول run
        if para.runs:
            para.runs[0].text = new_full
            for r in para.runs[1:]:
                r.text = ""


def fill_tasjil_wasl(mtrsh, markaz, montant_paye,
                     montant_paye_precedent=0.0, output_dir=None):
    """
    ملء نموذج Word (شهادة التسجيل + وصل التسديد) ببيانات المتربص.

    المعاملات:
        mtrsh                : قاموس بيانات المتربص
        markaz               : قاموس بيانات المركز
        montant_paye         : المبلغ المسدَّد حالياً
        montant_paye_precedent: المبلغ المسدَّد سابقاً
        output_dir           : مجلد الحفظ
    الإرجاع:
        مسار ملف Word الناتج
    """
    from docx import Document
    from docx.shared import Cm

    if not os.path.exists(TEMPLATE_FILE):
        raise FileNotFoundError(f"النموذج غير موجود: {TEMPLATE_FILE}")

    out_dir = output_dir or OUTPUT_DIR
    os.makedirs(out_dir, exist_ok=True)

    takh       = clean_takh(str(mtrsh.get("تخصص التكوين", "") or ""))
    ville      = extract_ville(markaz)
    today      = datetime.now().strftime("%Y/%m/%d")
    total      = get_montant(takh)

    # شخص طبيعي: ممثل المركز = اسم المركز تلقائياً
    # شخص معنوي: ممثل المركز مختلف عن اسم المركز
    representant = str(markaz.get("ممثل مركز التكوين", "") or "").strip()
    if not representant:
        representant = str(markaz.get("اسم مركز التكوين", "") or "")

    paye_avant = float(montant_paye_precedent or 0)
    paye_now   = float(montant_paye or 0)
    total_paye = paye_avant + paye_now
    reste      = total - total_paye
    tasjil_d   = fmt_date(str(mtrsh.get("تاريخ التسجيل", "") or today))

    # رقم الملف
    raqm_milaf = str(mtrsh.get("رقم الملف", "") or "")

    # ── جدول الاستبدال — الأطول أولاً لتجنب التداخل ──────────
    replacements = {
        # شهادة التسجيل — الأطول أولاً
        "1818": raqm_milaf,
        "1717": today,
        "1616": ville,
        "1515": "‏" + tasjil_d,
        "1414": fmt_raqm(str(mtrsh.get("رقم الدورة", "") or "")),
        "1313": takh,
        "1212": str(mtrsh.get("البريد الإلكتروني", "") or ""),
        "1111": str(mtrsh.get("رقم هاتف المتربص", "") or ""),
        "1010": str(mtrsh.get("العنوان", "") or ""),
        "999":  fmt_date(str(mtrsh.get("تاريخ صدور رخصة السياقة", "") or "")),
        "888":  str(mtrsh.get("رقم رخصة السياقة", "") or ""),
        "777":  str(mtrsh.get("رقم التعريف الوطني", "") or ""),
        "666":  str(mtrsh.get("تاريخ ومكان الازدياد", "") or ""),
        "555":  str(mtrsh.get("الاسم واللقب", "") or ""),
        "444":  str(markaz.get("البريد الإلكتروني", "") or ""),
        "333":  str(markaz.get("رقم الهاتف", "") or ""),
        "222":  representant,
        "111":  str(markaz.get("اسم مركز التكوين", "") or ""),
        # وصل التسديد — الأطول أولاً
        "3131": raqm_milaf,
        "3030": today,
        "2929": ville,
        "2828": f"{reste:,.2f}",
        "2727": f"{paye_now:,.2f}",
        "2626": f"{total:,.2f}",
        "2525": fmt_raqm(str(mtrsh.get("رقم الدورة", "") or "")),
        "2424": takh,
        "2323": str(mtrsh.get("الاسم واللقب", "") or ""),
    }

    doc = Document(TEMPLATE_FILE)

    # ── QR أولاً قبل أي استبدال (لتجنب تداخل الأرقام) ─────────
    qr1_text = (
        f"نوع الوثيقة: شهادة تسجيل\n"
        f"رقم الملف: {mtrsh.get('رقم الملف','')}\n"
        f"الاسم: {mtrsh.get('الاسم واللقب','')}\n"
        f"رقم التعريف: {mtrsh.get('رقم التعريف الوطني','')}\n"
        f"التخصص: {takh}\n"
        f"الدورة: {mtrsh.get('رقم الدورة','')}\n"
        f"تاريخ التسجيل: {tasjil_d}\n"
        f"المركز: {markaz.get('اسم مركز التكوين','')}"
    )
    qr2_text = (
        f"نوع الوثيقة: وصل تسديد\n"
        f"الاسم: {mtrsh.get('الاسم واللقب','')}\n"
        f"التخصص: {takh}\n"
        f"الدورة: {mtrsh.get('رقم الدورة','')}\n"
        f"الإجمالي: {total:,.2f} دج\n"
        f"المسدَّد: {total_paye:,.2f} دج\n"
        f"الباقي: {reste:,.2f} دج\n"
        f"التاريخ: {today}"
    )

    qr_map = {"9999": qr1_text, "88888": qr2_text}

    def insert_qr_in_para(para, key, qr_text):
        full_text = "".join(r.text or "" for r in para.runs)
        if key not in full_text:
            return False
        p = para._p
        for r in para.runs:
            p.remove(r._r)
        before = full_text.split(key)[0]
        if before:
            para.add_run(before)
        qr_run = para.add_run()
        qr_run.add_picture(make_qr(qr_text), width=Cm(2.5))
        after = full_text.split(key, 1)[1]
        if after:
            para.add_run(after)
        return True

    for para in doc.paragraphs:
        for key, qr_text in qr_map.items():
            if key in "".join(r.text or "" for r in para.runs):
                insert_qr_in_para(para, key, qr_text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for key, qr_text in qr_map.items():
                        if key in "".join(r.text or "" for r in para.runs):
                            insert_qr_in_para(para, key, qr_text)

    # ── استبدال النص بعد QR ────────────────────────────────────
    for para in doc.paragraphs:
        replace_in_para(para, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_para(para, replacements)

    # ── حفظ الملف ─────────────────────────────────────────────
    ts  = datetime.now().strftime("%Y%m%d_%H%M%S")
    nom = str(mtrsh.get("الاسم واللقب", "متربص")).replace(" ", "_")
    out = os.path.join(out_dir, f"تسجيل_وصل_{str(nom).replace("/","-")}_{ts}.docx")
    fix_all_dates_in_doc(doc)
    doc.save(out)
    print(f"✅ Word → {out}")
    return out


TEMPLATE_AQD = os.path.join(TEMPLATE_DIR, "عقد_التكوين.docx")


def fill_aqd_takwin(mtrsh, markaz, raqm_aqd, output_dir=None):
    """
    ملء نموذج Word لعقد التكوين ببيانات المتربص.

    المعاملات:
        mtrsh    : قاموس بيانات المتربص
        markaz   : قاموس بيانات المركز
        raqm_aqd : رقم العقد التسلسلي
        output_dir: مجلد الحفظ
    """
    from docx import Document
    from docx.shared import Cm

    if not os.path.exists(TEMPLATE_AQD):
        raise FileNotFoundError(f"النموذج غير موجود: {TEMPLATE_AQD}")

    out_dir = output_dir or OUTPUT_DIR
    os.makedirs(out_dir, exist_ok=True)

    takh   = clean_takh(str(mtrsh.get("تخصص التكوين", "") or ""))
    ville  = extract_ville(markaz)
    today  = datetime.now().strftime("%Y/%m/%d")
    total  = get_montant(takh)

    representant = str(markaz.get("ممثل مركز التكوين", "") or "").strip()
    if not representant:
        representant = str(markaz.get("اسم مركز التكوين", "") or "")

    replacements = {
        "4111": str(markaz.get("اسم مركز التكوين", "") or ""),
        "4222": representant,
        "4333": str(markaz.get("رقم الهاتف", "") or ""),
        "4334": str(markaz.get("البريد الإلكتروني", "") or ""),
        "4444": str(markaz.get("عنوان مركز التكوين", "") or ""),
        "4555": str(markaz.get("رقم الاعتماد", "") or ""),
        "4666": str(mtrsh.get("الاسم واللقب", "") or ""),
        "4777": str(mtrsh.get("تاريخ ومكان الازدياد", "") or ""),
        "4888": str(mtrsh.get("رقم التعريف الوطني", "") or ""),
        "4999": str(mtrsh.get("العنوان", "") or ""),
        "5111": str(mtrsh.get("رقم هاتف المتربص", "") or ""),
        "5222": takh,
        "5333": str(mtrsh.get("رقم الدورة", "") or ""),
        "5444": f"{total:,.2f}",
        "5555": ville,
        "5666": today,
        "[رقم تسلسلي]": str(raqm_aqd),
    }

    doc = Document(TEMPLATE_AQD)

    # ── QR أولاً ─────────────────────────────────────────────
    qr_text = (
        f"نوع الوثيقة: عقد تكوين\n"
        f"رقم العقد: {raqm_aqd}\n"
        f"الاسم: {mtrsh.get('الاسم واللقب','')}\n"
        f"رقم التعريف: {mtrsh.get('رقم التعريف الوطني','')}\n"
        f"التخصص: {takh}\n"
        f"الدورة: {mtrsh.get('رقم الدورة','')}\n"
        f"المبلغ: {total:,.2f} دج\n"
        f"المركز: {markaz.get('اسم مركز التكوين','')}\n"
        f"التاريخ: {today}"
    )

    def insert_qr_in_para(para, key, qr_text):
        full_text = "".join(r.text or "" for r in para.runs)
        if key not in full_text:
            return False
        p = para._p
        for r in para.runs:
            p.remove(r._r)
        before = full_text.split(key)[0]
        if before:
            para.add_run(before)
        para.add_run().add_picture(make_qr(qr_text), width=Cm(2.5))
        after = full_text.split(key, 1)[1]
        if after:
            para.add_run(after)
        return True

    # ── QR — استبدال 7777 بصورة مباشرة في مكانه ─────────────
    for para in doc.paragraphs:
        if "7777" in "".join(r.text or "" for r in para.runs):
            insert_qr_in_para(para, "7777", qr_text)
            break

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if "7777" in "".join(r.text or "" for r in para.runs):
                        insert_qr_in_para(para, "7777", qr_text)

    # ── استبدال النص ──────────────────────────────────────────
    for para in doc.paragraphs:
        replace_in_para(para, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_para(para, replacements)

    # ── حفظ الملف ─────────────────────────────────────────────
    ts  = datetime.now().strftime("%Y%m%d_%H%M%S")
    nom = str(mtrsh.get("الاسم واللقب", "متربص")).replace(" ", "_")
    raqm_safe = str(raqm_aqd).replace("/", "_")
    out = os.path.join(out_dir, f"عقد_تكوين_{raqm_safe}_{nom}_{ts}.docx")
    fix_all_dates_in_doc(doc)
    doc.save(out)
    print(f"✅ عقد Word → {out}")
    return out


def open_word(path):
    """فتح ملف Word بالتطبيق الافتراضي"""
    if not path or not os.path.exists(path): return
    if   platform.system() == "Windows": os.startfile(path)
    elif platform.system() == "Darwin":  subprocess.Popen(["open", path])
    else:                                subprocess.Popen(["xdg-open", path])


if __name__ == "__main__":
    # اختبار
    from طباعة_الشهادات import load_markaz, load_najihin
    markaz  = load_markaz()
    najihin = load_najihin()
    if najihin:
        path = fill_tasjil_wasl(najihin[0], markaz, montant_paye=8500)
        if path: open_word(path)
