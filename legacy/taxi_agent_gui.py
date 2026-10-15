#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
وكيل إدخال بيانات سيارات الأجرة - واجهة رسومية
انقر مرتين لفتح البرنامج
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import threading
import sys
import os
import json
import base64
import shutil
import datetime
import re
from pathlib import Path

# ═══════════════════════════════════════════════════════════════
#   إعدادات - عدّلها حسب جهازك
# ═══════════════════════════════════════════════════════════════
DEFAULT_API_KEY   = ""   # ← ضع مفتاح API هنا  sk-ant-api03-...
DEFAULT_EXCEL     = r"F:\Documents\تحديث ملف سيارات الاجرة\TAXI 32000.xlsx"
MODEL             = "claude-opus-4-5"
# ═══════════════════════════════════════════════════════════════


EXTRACTION_PROMPT = """
أنت خبير في قراءة وثائق رخص الاستغلال الجزائرية لسيارات الأجرة.
استخرج جميع المعلومات من هذه الصور وأعد JSON فقط بهذا الهيكل (بدون أي نص إضافي):

{
  "operateur": {
    "id_wilaya": 32,
    "nin": "رقم التعريف الوطني لصاحب/ة الرخصة 18 رقم",
    "prenom": "الاسم بالفرنسية",
    "nom": "اللقب بالفرنسية",
    "prenom_ar": "الاسم بالعربية",
    "nom_ar": "اللقب بالعربية",
    "date_naissance": "YYYY-MM-DD",
    "sexe": "M أو F",
    "adresse": "العنوان بالفرنسية",
    "adresse_ar": "العنوان بالعربية",
    "id_wilaya_adresse": 32,
    "n-porte": "رقم الباب",
    "date_debut_contrat": "YYYY-MM-DD",
    "date_fin_contrat": "YYYY-MM-DD",
    "qualite_oper_physique": "صفة المستفيد: أرملة مجاهد / ابن شهيد / ابنة شهيد",
    "activite": "Transports par Taxi",
    "nature_activite": "Individuelle أو inter-wilayas",
    "numero_casnos": null,
    "numero_cnas": null
  },
  "chauffeur": {
    "id_wilaya": 32,
    "activite": "Transports par Taxi",
    "nature_activite": "Individuelle أو inter-wilayas",
    "qualite_chauffeur": "مستفيد أو مكتري",
    "nin": "NIN السائق 18 رقم",
    "prenom": "اسم السائق بالفرنسية",
    "nom": "لقب السائق بالفرنسية",
    "prenom_ar": "اسم السائق بالعربية",
    "nom_ar": "لقب السائق بالعربية",
    "date_naissance": "YYYY-MM-DD",
    "sexe": "M",
    "adresse": "عنوان السائق بالفرنسية",
    "adresse_ar": "عنوان السائق بالعربية",
    "id_wilaya_adresse": 32,
    "numero_cnas": null,
    "numero_casnos": null,
    "n_permis": "رقم رخصة السياقة",
    "categorie_permis": "B أو D",
    "n_certificat_formation": "رقم شهادة التكوين",
    "date_certificat_formation": "YYYY-MM-DD",
    "n_telephone": null
  },
  "vehicule": {
    "id_wilaya": 32,
    "activite": "Transports par Taxi",
    "nature_activite": "Individuelle أو inter-wilayas",
    "n_porte": "رقم الباب",
    "n_immatriculation": "رقم التسجيل",
    "n_chasis": "الرقم التسلسلي VIN",
    "marque": "الصنف RENAULT/TOYOTA إلخ",
    "type": "الطراز",
    "n_siege": "عدد المقاعد",
    "energie": "ESSENCE أو GASOIL أو GPL",
    "annee_premier_utilisation": "السنة",
    "n_certificat_cerculation": "رقم رخصة السير",
    "date_fin_validite_certificat": "YYYY-MM-DD"
  },
  "licence_modjahid": {
    "id_wilaya": 32,
    "n_porte": "رقم الباب",
    "nin_modjahid": null,
    "prenom": "اسم المجاهد/الشهيد بالفرنسية",
    "nom": "لقب المجاهد/الشهيد بالفرنسية",
    "prenom_ar": "اسم المجاهد/الشهيد بالعربية",
    "nom_ar": "لقب المجاهد/الشهيد بالعربية",
    "date_naissance": "YYYY-MM-DD",
    "type_licence": "نوع الرخصة",
    "n_decision_octroi": "رقم قرار المنح",
    "date_decision_octroi": "YYYY-MM-DD"
  },
  "notes": "أي ملاحظات إضافية"
}
"""


def check_and_install(pkg, import_name=None):
    import importlib
    name = import_name or pkg
    try:
        importlib.import_module(name)
        return True
    except ImportError:
        return False


class TaxiAgentApp:
    def __init__(self, root):
        self.root = root
        self.root.title("وكيل سيارات الأجرة - ولاية البيض 32")
        self.root.geometry("750x650")
        self.root.resizable(True, True)
        self.root.configure(bg="#f0f4f8")

        # ── متغيرات ──
        self.pdf_path   = tk.StringVar()
        self.excel_path = tk.StringVar(value=DEFAULT_EXCEL)
        self.api_key    = tk.StringVar(value=DEFAULT_API_KEY)
        self.extracted  = None
        self.running    = False

        self._build_ui()

    def _build_ui(self):
        BG   = "#f0f4f8"
        BLUE = "#1a56db"
        WHITE= "#ffffff"
        FONT_AR = ("Arial", 11)

        # ── عنوان ──
        header = tk.Frame(self.root, bg=BLUE, pady=12)
        header.pack(fill="x")
        tk.Label(header, text="🚕  وكيل إدخال بيانات سيارات الأجرة",
                 font=("Arial", 16, "bold"), fg=WHITE, bg=BLUE).pack()
        tk.Label(header, text="ولاية البيض — رقم 32",
                 font=("Arial", 10), fg="#c7d9f5", bg=BLUE).pack()

        # ── حاوية رئيسية ──
        main = tk.Frame(self.root, bg=BG, padx=20, pady=15)
        main.pack(fill="both", expand=True)

        def section(parent, title):
            f = tk.LabelFrame(parent, text=f"  {title}  ",
                              font=("Arial", 10, "bold"), bg=BG,
                              fg="#374151", relief="groove", bd=2)
            f.pack(fill="x", pady=6)
            return f

        # ── مفتاح API ──
        sec1 = section(main, "🔑  مفتاح API")
        tk.Label(sec1, text="مفتاح Anthropic API:", font=FONT_AR, bg=BG).grid(row=0, column=0, sticky="e", padx=8, pady=6)
        e_api = tk.Entry(sec1, textvariable=self.api_key, width=52, show="*", font=("Courier", 10))
        e_api.grid(row=0, column=1, sticky="w", padx=4)
        tk.Button(sec1, text="إظهار", font=("Arial", 9),
                  command=lambda: e_api.config(show="" if e_api.cget("show") == "*" else "*")
                  ).grid(row=0, column=2, padx=4)

        # ── ملف PDF ──
        sec2 = section(main, "📄  ملف PDF للسائق")
        tk.Label(sec2, text="مسار الملف:", font=FONT_AR, bg=BG).grid(row=0, column=0, sticky="e", padx=8, pady=6)
        tk.Entry(sec2, textvariable=self.pdf_path, width=48, font=("Arial", 10)).grid(row=0, column=1, sticky="w", padx=4)
        tk.Button(sec2, text="اختر...", font=("Arial", 9), bg="#e5e7eb",
                  command=self.browse_pdf).grid(row=0, column=2, padx=4)

        # ── ملف Excel ──
        sec3 = section(main, "📊  ملف Excel")
        tk.Label(sec3, text="مسار Excel:", font=FONT_AR, bg=BG).grid(row=0, column=0, sticky="e", padx=8, pady=6)
        tk.Entry(sec3, textvariable=self.excel_path, width=48, font=("Arial", 10)).grid(row=0, column=1, sticky="w", padx=4)
        tk.Button(sec3, text="اختر...", font=("Arial", 9), bg="#e5e7eb",
                  command=self.browse_excel).grid(row=0, column=2, padx=4)

        # ── أزرار ──
        btn_frame = tk.Frame(main, bg=BG)
        btn_frame.pack(pady=10)
        self.btn_extract = tk.Button(btn_frame, text="🔍  استخراج البيانات",
                                      font=("Arial", 12, "bold"), bg=BLUE, fg=WHITE,
                                      padx=20, pady=8, relief="flat", cursor="hand2",
                                      command=self.start_extraction)
        self.btn_extract.pack(side="left", padx=8)

        self.btn_save = tk.Button(btn_frame, text="💾  حفظ في Excel",
                                   font=("Arial", 12, "bold"), bg="#059669", fg=WHITE,
                                   padx=20, pady=8, relief="flat", cursor="hand2",
                                   state="disabled", command=self.save_to_excel)
        self.btn_save.pack(side="left", padx=8)

        tk.Button(btn_frame, text="🗑️  مسح", font=("Arial", 11), bg="#e5e7eb",
                  padx=12, pady=8, relief="flat", command=self.clear_log).pack(side="left", padx=8)

        # ── شريط التقدم ──
        self.progress = ttk.Progressbar(main, mode="indeterminate", length=400)
        self.progress.pack(pady=4)

        # ── سجل العمليات ──
        sec4 = section(main, "📋  سجل العمليات")
        self.log = scrolledtext.ScrolledText(sec4, height=14, font=("Courier New", 9),
                                              bg="#1e293b", fg="#e2e8f0",
                                              insertbackground="white", wrap="word")
        self.log.pack(fill="both", expand=True, padx=6, pady=6)

        # ── شريط الحالة ──
        self.status_var = tk.StringVar(value="جاهز")
        tk.Label(self.root, textvariable=self.status_var, font=("Arial", 9),
                 bg="#e5e7eb", fg="#374151", anchor="w", relief="sunken"
                 ).pack(fill="x", side="bottom")

        self.log_msg("✅ البرنامج جاهز. اختر ملف PDF ثم اضغط 'استخراج البيانات'.\n")

    def browse_pdf(self):
        path = filedialog.askopenfilename(
            title="اختر ملف PDF",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")]
        )
        if path:
            self.pdf_path.set(path)
            self.log_msg(f"📂 تم اختيار: {Path(path).name}\n")

    def browse_excel(self):
        path = filedialog.askopenfilename(
            title="اختر ملف Excel",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if path:
            self.excel_path.set(path)

    def log_msg(self, msg):
        self.log.insert("end", msg)
        self.log.see("end")
        self.root.update_idletasks()

    def clear_log(self):
        self.log.delete("1.0", "end")
        self.extracted = None
        self.btn_save.config(state="disabled")

    def start_extraction(self):
        if self.running:
            return
        if not self.pdf_path.get():
            messagebox.showwarning("تنبيه", "الرجاء اختيار ملف PDF أولاً")
            return
        if not self.api_key.get().strip():
            messagebox.showwarning("تنبيه", "الرجاء إدخال مفتاح API")
            return
        if not os.path.exists(self.pdf_path.get()):
            messagebox.showerror("خطأ", "ملف PDF غير موجود!")
            return

        self.running = True
        self.btn_extract.config(state="disabled")
        self.btn_save.config(state="disabled")
        self.progress.start(10)
        self.status_var.set("جارٍ الاستخراج...")

        t = threading.Thread(target=self._run_extraction, daemon=True)
        t.start()

    def _run_extraction(self):
        try:
            # فحص المكتبات
            missing = []
            for pkg, imp in [("anthropic", "anthropic"), ("pymupdf", "fitz"), ("openpyxl", "openpyxl")]:
                if not check_and_install(pkg, imp):
                    missing.append(pkg)
            if missing:
                self.log_msg(f"❌ مكتبات مفقودة: {', '.join(missing)}\n")
                self.log_msg(f"   شغّل في CMD: pip install {' '.join(missing)}\n")
                return

            import anthropic
            import fitz

            # تحويل PDF إلى صور
            self.log_msg("🔄 تحويل PDF إلى صور...\n")
            pdf_path = self.pdf_path.get()
            doc = fitz.open(pdf_path)
            self.log_msg(f"   عدد الصفحات: {len(doc)}\n")
            images = []
            for i, page in enumerate(doc):
                mat = fitz.Matrix(2.0, 2.0)
                pix = page.get_pixmap(matrix=mat)
                img_data = pix.tobytes("jpeg")
                b64 = base64.standard_b64encode(img_data).decode("utf-8")
                images.append({
                    "type": "image",
                    "source": {"type": "base64", "media_type": "image/jpeg", "data": b64}
                })
                self.log_msg(f"   ✓ صفحة {i+1}/{len(doc)}\n")
            doc.close()

            # إرسال إلى Claude
            self.log_msg(f"\n🤖 إرسال إلى Claude للاستخراج...\n")
            client = anthropic.Anthropic(api_key=self.api_key.get().strip())
            content = images + [{"type": "text", "text": EXTRACTION_PROMPT}]
            message = client.messages.create(
                model=MODEL,
                max_tokens=4096,
                messages=[{"role": "user", "content": content}]
            )
            raw = message.content[0].text.strip()

            # تحليل JSON
            json_match = re.search(r'\{.*\}', raw, re.DOTALL)
            if json_match:
                raw = json_match.group(0)
            data = json.loads(raw)
            self.extracted = data

            # عرض النتائج
            op = data.get("operateur", {})
            ch = data.get("chauffeur", {})
            ve = data.get("vehicule", {})
            lm = data.get("licence_modjahid", {})

            self.log_msg("\n" + "═"*55 + "\n")
            self.log_msg("✅ تم الاستخراج بنجاح!\n\n")
            self.log_msg(f"🏠 صاحب/ة الرخصة:\n")
            self.log_msg(f"   {op.get('prenom_ar','')} {op.get('nom_ar','')} | NIN: {op.get('nin','')}\n")
            self.log_msg(f"   رقم الباب: {op.get('n-porte','')} | الرخصة: {op.get('date_debut_contrat','')} → {op.get('date_fin_contrat','')}\n")
            self.log_msg(f"   الصفة: {op.get('qualite_oper_physique','')}\n\n")

            self.log_msg(f"🚗 السائق:\n")
            self.log_msg(f"   {ch.get('prenom_ar','')} {ch.get('nom_ar','')} | NIN: {ch.get('nin','')}\n")
            self.log_msg(f"   الصفة: {ch.get('qualite_chauffeur','')} | رخصة القيادة: {ch.get('n_permis','')}\n\n")

            self.log_msg(f"🚙 المركبة:\n")
            self.log_msg(f"   {ve.get('marque','')} {ve.get('type','')} | التسجيل: {ve.get('n_immatriculation','')}\n")
            self.log_msg(f"   الوقود: {ve.get('energie','')} | المقاعد: {ve.get('n_siege','')} | السنة: {ve.get('annee_premier_utilisation','')}\n\n")

            self.log_msg(f"📜 رخصة المجاهد:\n")
            self.log_msg(f"   {lm.get('prenom_ar','')} {lm.get('nom_ar','')} | {lm.get('type_licence','')}\n")
            self.log_msg(f"   قرار المنح: {lm.get('n_decision_octroi','')} بتاريخ {lm.get('date_decision_octroi','')}\n")
            self.log_msg("═"*55 + "\n\n")
            self.log_msg("⬆️  اضغط 'حفظ في Excel' لإضافة البيانات إلى الملف\n")

            # حفظ JSON
            json_out = pdf_path.replace(".pdf", "_extracted.json")
            with open(json_out, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            self.log_msg(f"💾 JSON محفوظ: {Path(json_out).name}\n")

            self.btn_save.config(state="normal")
            self.status_var.set("✅ تم الاستخراج — اضغط 'حفظ في Excel'")

        except Exception as e:
            self.log_msg(f"\n❌ خطأ: {e}\n")
            self.status_var.set(f"خطأ: {e}")
        finally:
            self.running = False
            self.btn_extract.config(state="normal")
            self.progress.stop()

    def save_to_excel(self):
        if not self.extracted:
            return
        excel_path = self.excel_path.get()
        if not os.path.exists(excel_path):
            messagebox.showerror("خطأ", f"ملف Excel غير موجود:\n{excel_path}")
            return

        if not messagebox.askyesno("تأكيد", "هل تريد إضافة البيانات إلى ملف Excel؟"):
            return

        try:
            import openpyxl

            # نسخة احتياطية
            backup = excel_path.replace(".xlsx", f"_backup_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            shutil.copy2(excel_path, backup)
            self.log_msg(f"\n💾 نسخة احتياطية: {Path(backup).name}\n")

            wb = openpyxl.load_workbook(excel_path)
            data = self.extracted

            n_porte = str(
                data.get("operateur", {}).get("n-porte") or
                data.get("vehicule", {}).get("n_porte") or
                data.get("licence_modjahid", {}).get("n_porte") or ""
            )

            sheets_config = [
                ("Operateur",        data.get("operateur", {}),        "n-porte",   n_porte),
                ("Chauffeur",        data.get("chauffeur", {}),         None,        None),
                ("Véhicule",         data.get("vehicule", {}),          "n_porte",   n_porte),
                ("Licence Modjahid", data.get("licence_modjahid", {}),  "n_porte",   n_porte),
            ]

            for sheet_name, row_data, dup_field, dup_val in sheets_config:
                if not row_data or sheet_name not in wb.sheetnames:
                    continue
                ws = wb[sheet_name]

                # خريطة الأعمدة
                col_map = {}
                for ci, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)), 1):
                    if cell is not None:
                        col_map[str(cell).strip()] = ci

                # فحص التكرار
                if dup_field and dup_val:
                    dup_ci = col_map.get(dup_field)
                    if dup_ci:
                        for ri in range(3, ws.max_row + 1):
                            v = ws.cell(row=ri, column=dup_ci).value
                            if v is not None and str(v).strip() == dup_val:
                                self.log_msg(f"⚠️  [{sheet_name}] رقم الباب {dup_val} موجود — تم التخطي\n")
                                break
                        else:
                            # إضافة الصف
                            row_idx = self._next_empty(ws)
                            self._write_row(ws, col_map, row_data, row_idx)
                            self.log_msg(f"✅ [{sheet_name}] تمت الإضافة في الصف {row_idx}\n")
                        continue

                row_idx = self._next_empty(ws)
                self._write_row(ws, col_map, row_data, row_idx)
                self.log_msg(f"✅ [{sheet_name}] تمت الإضافة في الصف {row_idx}\n")

            wb.save(excel_path)
            self.log_msg(f"\n🎉 تم الحفظ في: {Path(excel_path).name}\n")
            self.status_var.set("✅ تم الحفظ بنجاح!")
            messagebox.showinfo("تم", "✅ تمت إضافة البيانات بنجاح في ملف Excel!")
            self.btn_save.config(state="disabled")

        except Exception as e:
            self.log_msg(f"\n❌ خطأ أثناء الحفظ: {e}\n")
            messagebox.showerror("خطأ", str(e))

    def _next_empty(self, ws) -> int:
        for ri in range(3, ws.max_row + 2):
            if ws.cell(row=ri, column=1).value is None:
                return ri
        return ws.max_row + 1

    def _write_row(self, ws, col_map: dict, data: dict, row_idx: int):
        for field, value in data.items():
            if value is None:
                continue
            col_idx = None
            # بحث مرن في اسم العمود
            normalized = field.replace("-", "_").lower()
            for key, ci in col_map.items():
                if key.replace("-", "_").lower() == normalized:
                    col_idx = ci
                    break
            if col_idx:
                ws.cell(row=row_idx, column=col_idx).value = value


def main():
    # فحص Python
    if sys.version_info < (3, 8):
        import tkinter.messagebox as mb
        mb.showerror("خطأ", "يتطلب البرنامج Python 3.8 أو أحدث")
        sys.exit(1)

    root = tk.Tk()
    app = TaxiAgentApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
