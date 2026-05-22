#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
وحدة قاعدة البيانات SQLite
- تخزين داخلي سريع
- نسخ احتياطي تلقائي
- تشفير البيانات الحساسة
- تصدير لـ Excel عند الحاجة
"""
import os, sqlite3, shutil, hashlib, base64
from datetime import datetime
from pathlib import Path

DB_DIR  = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(DB_DIR, "نظام_الكفاءة.db")
BACKUP_DIR = os.path.join(DB_DIR, "backup")

# ── مفتاح التشفير ────────────────────────────────────────────
def _get_key():
    """مفتاح تشفير مبني على اسم الجهاز"""
    import socket
    machine = socket.gethostname()
    return hashlib.sha256(machine.encode()).hexdigest()[:32].encode()

def _encrypt(text):
    """تشفير بسيط للبيانات الحساسة"""
    if not text: return text
    key = _get_key()
    data = str(text).encode()
    enc  = bytes(b ^ key[i % len(key)] for i, b in enumerate(data))
    return base64.b64encode(enc).decode()

def _decrypt(text):
    """فك التشفير"""
    if not text: return text
    try:
        key  = _get_key()
        data = base64.b64decode(text.encode())
        dec  = bytes(b ^ key[i % len(key)] for i, b in enumerate(data))
        return dec.decode()
    except Exception:
        return text  # إذا كان النص غير مشفر نُرجعه كما هو

# ── الحقول الحساسة التي تُشفَّر ──────────────────────────────
SENSITIVE = {"رقم التعريف الوطني", "رقم رخصة السياقة", "رقم هاتف المتربص"}

# ══════════════════════════════════════════════════════════════
# إنشاء قاعدة البيانات
# ══════════════════════════════════════════════════════════════
def init_db():
    """إنشاء الجداول إذا لم تكن موجودة"""
    conn = sqlite3.connect(DB_FILE)
    c    = conn.cursor()

    # جدول المركز
    c.execute('''CREATE TABLE IF NOT EXISTS مركز (
        id INTEGER PRIMARY KEY,
        اسم_مركز_التكوين TEXT,
        رقم_الاعتماد TEXT,
        تاريخ_الاعتماد TEXT,
        عدد_الأقسام TEXT,
        مديرية_النقل_لولاية TEXT,
        عنوان_مركز_التكوين TEXT,
        بلدية_المركز TEXT,
        ممثل_مركز_التكوين TEXT,
        المسؤول_البيداغوجي TEXT,
        الأساتذة TEXT,
        رقم_الهاتف TEXT,
        البريد_الإلكتروني TEXT,
        آخر_رقم_ملف INTEGER DEFAULT 0,
        تاريخ_التحديث TEXT
    )''')
    # إضافة عمود بلدية_المركز إذا لم يكن موجوداً (للقواعد القديمة)
    try:
        c.execute("ALTER TABLE مركز ADD COLUMN بلدية_المركز TEXT DEFAULT ''")
    except Exception:
        pass

    # جدول الدورات
    c.execute('''CREATE TABLE IF NOT EXISTS دورات (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        رقم_الدورة TEXT UNIQUE NOT NULL,
        تاريخ_افتتاح_الدورة TEXT,
        تاريخ_اختتام_الدورة TEXT,
        عدد_المترشحين INTEGER DEFAULT 0,
        رقم_المحضر TEXT,
        تاريخ_الإنشاء TEXT DEFAULT CURRENT_TIMESTAMP
    )''')

    # جدول المترشحين
    c.execute('''CREATE TABLE IF NOT EXISTS مترشحون (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        رقم_الملف TEXT UNIQUE,
        تاريخ_التسجيل TEXT,
        الاسم_واللقب TEXT NOT NULL,
        اللقب_بالفرنسية TEXT,
        الاسم_بالفرنسية TEXT,
        تاريخ_ومكان_الازدياد TEXT,
        العنوان TEXT,
        رقم_التعريف_الوطني TEXT,
        رقم_رخصة_السياقة TEXT,
        تاريخ_صدور_رخصة_السياقة TEXT,
        تاريخ_انتهاء_رخصة_السياقة TEXT,
        مكان_إصدار_الرخصة TEXT,
        فئات_الرخصة TEXT,
        زمرة_الدم TEXT,
        الجنس TEXT,
        رقم_هاتف_المتربص TEXT,
        البريد_الإلكتروني TEXT,
        تخصص_التكوين TEXT,
        رقم_الدورة TEXT,
        حالة_النجاح TEXT DEFAULT '',
        رقم_الشهادة TEXT DEFAULT '',
        تاريخ_الشهادة TEXT DEFAULT '',
        رقم_الشهادة_2 TEXT DEFAULT '',
        تاريخ_الشهادة_2 TEXT DEFAULT '',
        مسار_صورة_الرخصة TEXT DEFAULT '',
        مسار_صورة_الشهادة TEXT DEFAULT '',
        FOREIGN KEY (رقم_الدورة) REFERENCES دورات(رقم_الدورة)
    )''')

    # جدول النقاط
    c.execute('''CREATE TABLE IF NOT EXISTS نقاط (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        رقم_الملف TEXT,
        رقم_الدورة TEXT,
        الاسم_واللقب TEXT,
        التخصص TEXT,
        نقاط TEXT,  -- JSON
        المعدل REAL,
        القرار TEXT,
        FOREIGN KEY (رقم_الملف) REFERENCES مترشحون(رقم_الملف)
    )''')

    conn.commit()
    conn.close()
    print("✅ قاعدة البيانات جاهزة")


# ══════════════════════════════════════════════════════════════
# النسخ الاحتياطي التلقائي
# ══════════════════════════════════════════════════════════════
def backup_db():
    """نسخ احتياطي تلقائي للقاعدة والـ Excel"""
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")

    # نسخ SQLite
    if os.path.exists(DB_FILE):
        dest = os.path.join(BACKUP_DIR, f"نظام_الكفاءة_{ts}.db")
        shutil.copy2(DB_FILE, dest)

    # نسخ Excel
    excel = os.path.join(DB_DIR, "بيانات_مركز_التكوين.xlsx")
    if os.path.exists(excel):
        dest_xl = os.path.join(BACKUP_DIR, f"بيانات_مركز_{ts}.xlsx")
        shutil.copy2(excel, dest_xl)

    # حذف النسخ القديمة (أكثر من 30 يوم)
    _cleanup_backups(days=30)
    print(f"✅ نسخة احتياطية: {ts}")
    return ts


def _cleanup_backups(days=30):
    """حذف النسخ الاحتياطية الأقدم من X يوم"""
    if not os.path.exists(BACKUP_DIR): return
    import time
    limit = time.time() - (days * 86400)
    for f in Path(BACKUP_DIR).iterdir():
        if f.is_file() and f.stat().st_mtime < limit:
            f.unlink()


def get_backup_list():
    """قائمة النسخ الاحتياطية المتاحة"""
    if not os.path.exists(BACKUP_DIR): return []
    files = sorted(Path(BACKUP_DIR).glob("*.db"), reverse=True)
    return [f.name for f in files]


def restore_backup(filename):
    """استعادة نسخة احتياطية"""
    src = os.path.join(BACKUP_DIR, filename)
    if not os.path.exists(src): return False
    # نسخ احتياطي للحالي قبل الاستعادة
    backup_db()
    shutil.copy2(src, DB_FILE)
    print(f"✅ تم استعادة: {filename}")
    return True


# ══════════════════════════════════════════════════════════════
# عمليات المركز
# ══════════════════════════════════════════════════════════════
def save_markaz(data):
    """حفظ بيانات المركز"""
    conn = sqlite3.connect(DB_FILE)
    c    = conn.cursor()
    c.execute("SELECT id FROM مركز LIMIT 1")
    row = c.fetchone()
    now = datetime.now().strftime("%Y/%m/%d %H:%M:%S")

    vals = (
        data.get("اسم مركز التكوين",""),
        data.get("رقم الاعتماد",""),
        data.get("تاريخ الاعتماد",""),
        data.get("عدد الأقسام",""),
        data.get("مديرية النقل لولاية",""),
        data.get("عنوان مركز التكوين",""),
        data.get("بلدية المركز",""),
        data.get("ممثل مركز التكوين",""),
        data.get("المسؤول البيداغوجي",""),
        data.get("الأساتذة",""),
        data.get("رقم الهاتف",""),
        data.get("البريد الإلكتروني",""),
        int(data.get("آخر رقم ملف",0) or 0),
        now
    )

    if row:
        c.execute("""UPDATE مركز SET
            اسم_مركز_التكوين=?, رقم_الاعتماد=?, تاريخ_الاعتماد=?,
            عدد_الأقسام=?, مديرية_النقل_لولاية=?, عنوان_مركز_التكوين=?,
            بلدية_المركز=?,
            ممثل_مركز_التكوين=?, المسؤول_البيداغوجي=?, الأساتذة=?,
            رقم_الهاتف=?, البريد_الإلكتروني=?, آخر_رقم_ملف=?, تاريخ_التحديث=?
            WHERE id=1""", vals)
    else:
        c.execute("""INSERT INTO مركز VALUES (1,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", vals)

    conn.commit(); conn.close()
    backup_db()


def load_markaz():
    """تحميل بيانات المركز"""
    conn = sqlite3.connect(DB_FILE)
    c    = conn.cursor()
    c.execute("SELECT * FROM مركز LIMIT 1")
    row = c.fetchone()
    conn.close()
    if not row: return {}
    keys = ["id","اسم مركز التكوين","رقم الاعتماد","تاريخ الاعتماد",
            "عدد الأقسام","مديرية النقل لولاية","عنوان مركز التكوين",
            "بلدية المركز",
            "ممثل مركز التكوين","المسؤول البيداغوجي","الأساتذة",
            "رقم الهاتف","البريد الإلكتروني","آخر رقم ملف","تاريخ التحديث"]
    return dict(zip(keys, row))


# ══════════════════════════════════════════════════════════════
# عمليات الدورات
# ══════════════════════════════════════════════════════════════
def save_dawra(raqm, date_ouv, date_clo, raqm_mahdar=""):
    conn = sqlite3.connect(DB_FILE)
    c    = conn.cursor()
    # إضافة عمود رقم_المحضر إذا لم يكن موجوداً
    try:
        c.execute("ALTER TABLE دورات ADD COLUMN رقم_المحضر TEXT")
        conn.commit()
    except Exception: pass
    c.execute("""INSERT OR REPLACE INTO دورات
        (رقم_الدورة, تاريخ_افتتاح_الدورة, تاريخ_اختتام_الدورة, رقم_المحضر)
        VALUES (?,?,?,?)""", (raqm, date_ouv, date_clo, raqm_mahdar))
    conn.commit(); conn.close()
    backup_db()


def load_dawrat():
    conn = sqlite3.connect(DB_FILE)
    c    = conn.cursor()
    c.execute("SELECT رقم_الدورة, تاريخ_افتتاح_الدورة, تاريخ_اختتام_الدورة, عدد_المترشحين FROM دورات ORDER BY id")
    rows = c.fetchall()
    conn.close()
    return rows


def delete_dawra(raqm):
    conn = sqlite3.connect(DB_FILE)
    conn.cursor().execute("DELETE FROM دورات WHERE رقم_الدورة=?", (raqm,))
    conn.commit(); conn.close()


# ══════════════════════════════════════════════════════════════
# عمليات المترشحين
# ══════════════════════════════════════════════════════════════
def next_file_num_db():
    """توليد رقم الملف التالي من SQLite"""
    conn = sqlite3.connect(DB_FILE)
    c    = conn.cursor()
    c.execute("SELECT آخر_رقم_ملف FROM مركز LIMIT 1")
    row  = c.fetchone()
    last = int(row[0]) if row and row[0] else 0
    new  = last + 1
    c.execute("UPDATE مركز SET آخر_رقم_ملف=? WHERE id=1", (new,))
    conn.commit(); conn.close()
    return str(new).zfill(4)


def save_mtrsh(data):
    """حفظ مترشح مع تشفير البيانات الحساسة"""
    conn = sqlite3.connect(DB_FILE)
    c    = conn.cursor()

    # تشفير الحقول الحساسة
    nin   = _encrypt(data.get("رقم التعريف الوطني",""))
    rukh  = _encrypt(data.get("رقم رخصة السياقة",""))
    tel   = _encrypt(data.get("رقم هاتف المتربص",""))

    c.execute("""INSERT OR REPLACE INTO مترشحون VALUES (
        NULL,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?
    )""", (
        data.get("رقم الملف",""),
        data.get("تاريخ التسجيل",""),
        data.get("الاسم واللقب",""),
        data.get("اللقب بالفرنسية",""),
        data.get("الاسم بالفرنسية",""),
        data.get("تاريخ ومكان الازدياد",""),
        data.get("العنوان",""),
        nin, rukh,
        data.get("تاريخ صدور رخصة السياقة",""),
        data.get("تاريخ انتهاء رخصة السياقة",""),
        data.get("مكان إصدار الرخصة",""),
        data.get("فئات الرخصة",""),
        data.get("زمرة الدم",""),
        data.get("الجنس",""),
        tel,
        data.get("البريد الإلكتروني",""),
        data.get("تخصص التكوين",""),
        data.get("رقم الدورة",""),
        data.get("حالة النجاح",""),
        data.get("رقم الشهادة",""),
        data.get("تاريخ الشهادة",""),
        data.get("رقم الشهادة 2",""),
        data.get("تاريخ الشهادة 2",""),
        data.get("مسار صورة الرخصة",""),
        data.get("مسار صورة الشهادة",""),
    ))

    # تحديث عدد المترشحين في الدورة
    conn.cursor().execute("""UPDATE دورات SET عدد_المترشحين=(
        SELECT COUNT(*) FROM مترشحون WHERE رقم_الدورة=?
    ) WHERE رقم_الدورة=?""", (data.get("رقم الدورة",""),)*2)

    conn.commit(); conn.close()


def load_mtrsheen(raqm_dawra=None, search=None):
    """تحميل المترشحين مع فك التشفير"""
    conn = sqlite3.connect(DB_FILE)
    c    = conn.cursor()

    if raqm_dawra:
        c.execute("SELECT * FROM مترشحون WHERE رقم_الدورة=? ORDER BY id",
                  (str(raqm_dawra),))
    elif search:
        c.execute("SELECT * FROM مترشحون WHERE الاسم_واللقب LIKE ? ORDER BY id",
                  (f"%{search}%",))
    else:
        c.execute("SELECT * FROM مترشحون ORDER BY id")

    rows = c.fetchall()
    conn.close()

    cols = ["id","رقم الملف","تاريخ التسجيل","الاسم واللقب",
            "اللقب بالفرنسية","الاسم بالفرنسية","تاريخ ومكان الازدياد",
            "العنوان","رقم التعريف الوطني","رقم رخصة السياقة",
            "تاريخ صدور رخصة السياقة","تاريخ انتهاء رخصة السياقة",
            "مكان إصدار الرخصة","فئات الرخصة","زمرة الدم","الجنس",
            "رقم هاتف المتربص","البريد الإلكتروني","تخصص التكوين",
            "رقم الدورة","حالة النجاح","رقم الشهادة","تاريخ الشهادة",
            "رقم الشهادة 2","تاريخ الشهادة 2","مسار صورة الرخصة","مسار صورة الشهادة"]

    result = []
    for row in rows:
        d = dict(zip(cols, row))
        # فك تشفير الحقول الحساسة
        d["رقم التعريف الوطني"] = _decrypt(d["رقم التعريف الوطني"])
        d["رقم رخصة السياقة"]   = _decrypt(d["رقم رخصة السياقة"])
        d["رقم هاتف المتربص"]   = _decrypt(d["رقم هاتف المتربص"])
        result.append(d)
    return result


def update_mtrsh_status(raqm_milaf, status, num_sh="", date_sh="",
                         num_sh2="", date_sh2=""):
    """تحديث حالة النجاح والشهادة"""
    conn = sqlite3.connect(DB_FILE)
    conn.cursor().execute("""UPDATE مترشحون SET
        حالة_النجاح=?, رقم_الشهادة=?, تاريخ_الشهادة=?,
        رقم_الشهادة_2=?, تاريخ_الشهادة_2=?
        WHERE رقم_الملف=?""",
        (status, num_sh, date_sh, num_sh2, date_sh2, raqm_milaf))
    conn.commit(); conn.close()


def delete_mtrsh(raqm_milaf):
    conn = sqlite3.connect(DB_FILE)
    conn.cursor().execute("DELETE FROM مترشحون WHERE رقم_الملف=?", (raqm_milaf,))
    conn.commit(); conn.close()


# ══════════════════════════════════════════════════════════════
# تصدير لـ Excel
# ══════════════════════════════════════════════════════════════
def export_to_excel(output_path=None):
    """تصدير كل البيانات لـ Excel"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    output_path = output_path or os.path.join(DB_DIR, "بيانات_مركز_التكوين.xlsx")
    wb = openpyxl.Workbook()

    # ── ورقة المركز ──────────────────────────────────────────
    ws_m = wb.active; ws_m.title = "مركز_التكوين"
    m = load_markaz()
    keys = ["اسم مركز التكوين","رقم الاعتماد","تاريخ الاعتماد","عدد الأقسام",
            "مديرية النقل لولاية","عنوان مركز التكوين","ممثل مركز التكوين",
            "المسؤول البيداغوجي","الأساتذة","رقم الهاتف","البريد الإلكتروني","آخر رقم ملف"]
    ws_m.append(keys)
    ws_m.append([m.get(k,"") for k in keys])

    # ── ورقة الدورات ─────────────────────────────────────────
    ws_d = wb.create_sheet("الدورات")
    ws_d.append(["تاريخ افتتاح الدورة","تاريخ اختتام الدورة","رقم الدورة","العدد الإجمالي"])
    for r in load_dawrat():
        ws_d.append([r[1], r[2], r[0], r[3]])

    # ── ورقة المترشحين ───────────────────────────────────────
    ws_t = wb.create_sheet("المترشحون")
    cols = ["رقم الملف","تاريخ التسجيل","الاسم واللقب","اللقب بالفرنسية",
            "الاسم بالفرنسية","تاريخ ومكان الازدياد","العنوان",
            "رقم التعريف الوطني","رقم رخصة السياقة",
            "تاريخ صدور رخصة السياقة","تاريخ انتهاء رخصة السياقة",
            "مكان إصدار الرخصة","فئات الرخصة","زمرة الدم","الجنس",
            "رقم هاتف المتربص","البريد الإلكتروني","تخصص التكوين",
            "رقم الدورة","حالة النجاح","رقم الشهادة","تاريخ الشهادة",
            "رقم الشهادة 2","تاريخ الشهادة 2"]
    ws_t.append(cols)
    for m in load_mtrsheen():
        ws_t.append([m.get(c,"") for c in cols])

    wb.save(output_path)
    print(f"✅ تم التصدير → {output_path}")
    return output_path


# ══════════════════════════════════════════════════════════════
# استيراد من Excel القديم
# ══════════════════════════════════════════════════════════════
def import_from_excel(excel_path):
    """استيراد البيانات من Excel الموجود لـ SQLite"""
    import openpyxl
    wb = openpyxl.load_workbook(excel_path)

    imported = {"مترشحون": 0, "دورات": 0}

    # استيراد الدورات
    if "الدورات" in wb.sheetnames:
        ws = wb["الدورات"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row): continue
            raqm = str(row[2] or "").strip()
            if raqm:
                save_dawra(raqm, str(row[0] or ""), str(row[1] or ""))
                imported["دورات"] += 1

    # استيراد المترشحين
    if "المترشحون" in wb.sheetnames:
        ws = wb["المترشحون"]
        cols_old = ["تاريخ التسجيل","الاسم واللقب","اللقب بالفرنسية",
                    "الاسم بالفرنسية","تاريخ ومكان الازدياد","العنوان",
                    "رقم التعريف الوطني","رقم رخصة السياقة",
                    "تاريخ صدور رخصة السياقة","تاريخ انتهاء رخصة السياقة",
                    "مكان إصدار الرخصة","فئات الرخصة","زمرة الدم","الجنس",
                    "رقم هاتف المتربص","البريد الإلكتروني","تخصص التكوين",
                    "رقم الدورة","حالة النجاح","رقم الشهادة","تاريخ الشهادة",
                    "رقم الشهادة 2","تاريخ الشهادة 2","مسار صورة الرخصة",
                    "مسار صورة الشهادة","رقم الملف"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row): continue
            d = dict(zip(cols_old, list(row)+[""]*(len(cols_old)-len(row))))
            if d.get("الاسم واللقب"):
                if not d.get("رقم الملف"):
                    d["رقم الملف"] = next_file_num_db()
                save_mtrsh(d)
                imported["مترشحون"] += 1

    print(f"✅ استيراد: {imported}")
    return imported


if __name__ == "__main__":
    init_db()
    print("قاعدة البيانات جاهزة للاستخدام")
