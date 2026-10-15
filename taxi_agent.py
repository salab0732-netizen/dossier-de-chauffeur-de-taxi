#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=============================================================
  وكيل إدخال بيانات سيارات الأجرة - ولاية البيض (32)
  Taxi Data Entry Agent - El Bayadh Wilaya (32)
=============================================================
الاستخدام:
  python taxi_agent.py <ملف_PDF>  [مسار_اكسيل]

مثال:
  python taxi_agent.py "بلاكحل عبد القادر.pdf"
  python taxi_agent.py "محسين خالد.pdf" "TAXI 32000.xlsx"

المتطلبات:
  pip install anthropic pymupdf openpyxl
=============================================================
"""

import sys
import os
import json
import base64
import shutil
import datetime
import re
from pathlib import Path

# ─── إعدادات ─────────────────────────────────────────────────
API_KEY    = os.environ.get("ANTHROPIC_API_KEY", "")   # أو ضعه هنا مباشرة
MODEL      = "claude-opus-4-5"                           # أفضل نموذج للقراءة البصرية
WILAYA_ID  = 32                                          # ولاية البيض
EXCEL_PATH = r"F:\Documents\تحديث ملف سيارات الاجرة\TAXI 32000.xlsx"
# ─────────────────────────────────────────────────────────────


EXTRACTION_PROMPT = """
أنت خبير في قراءة وثائق رخص الاستغلال الجزائرية لسيارات الأجرة.
استخرج جميع المعلومات من هذه الصور (وثائق مسح ضوئي) وأعد JSON فقط بالهيكل التالي.

ملاحظات مهمة:
- الأرقام: اكتبها كما هي بدون تنسيق (مثال: "1290" لا "١٢٩٠")
- التواريخ: تنسيق YYYY-MM-DD
- إذا لم تجد المعلومة: اكتب null
- المستفيد (operateur) هو صاحب/صاحبة الرخصة (ذو حق المجاهد/الشهيد)
- السائق (chauffeur) هو من يقود التاكسي فعلياً (قد يكون نفس المستفيد أو مختلفاً)
- رقم الولاية دائماً 32 لولاية البيض

أعد JSON بهذا الهيكل بالضبط:

{
  "operateur": {
    "id_wilaya": 32,
    "nin": "رقم التعريف الوطني لصاحب/ة الرخصة (18 رقم)",
    "prenom": "الاسم بالفرنسية",
    "nom": "اللقب بالفرنسية",
    "prenom_ar": "الاسم بالعربية",
    "nom_ar": "اللقب بالعربية",
    "date_naissance": "YYYY-MM-DD",
    "id_commune_naissance": null,
    "id_wilaya_naissance": null,
    "sexe": "M أو F",
    "n_acte_naissance": "رقم شهادة الميلاد إن وُجد",
    "adresse": "العنوان بالفرنسية",
    "adresse_ar": "العنوان بالعربية",
    "id_wilaya_adresse": 32,
    "n-porte": "رقم الباب (رقم الرخصة على باب التاكسي)",
    "date_debut_contrat": "تاريخ بداية الرخصة YYYY-MM-DD",
    "date_fin_contrat": "تاريخ انتهاء الرخصة YYYY-MM-DD",
    "qualite_oper_physique": "ذوي حقوق (أرملة/ابن/ابنة شهيد أو مجاهد)",
    "activite": "Transports par Taxi",
    "nature_activite": "Individuelle أو inter-wilayas",
    "numero_casnos": "رقم CASNOS إن وُجد",
    "numero_cnas": "رقم CNAS إن وُجد"
  },
  "chauffeur": {
    "id_wilaya": 32,
    "activite": "Transports par Taxi",
    "nature_activite": "Individuelle أو inter-wilayas",
    "qualite_chauffeur": "مستفيد (يقود بنفسه) أو مكتري (مستأجر للرخصة)",
    "nin": "رقم التعريف الوطني للسائق (18 رقم)",
    "prenom": "اسم السائق بالفرنسية",
    "nom": "لقب السائق بالفرنسية",
    "prenom_ar": "اسم السائق بالعربية",
    "nom_ar": "لقب السائق بالعربية",
    "date_naissance": "YYYY-MM-DD",
    "id_wilaya_naissance": null,
    "sexe": "M",
    "adresse": "عنوان السائق بالفرنسية",
    "adresse_ar": "عنوان السائق بالعربية",
    "id_wilaya_adresse": 32,
    "numero_cnas": "رقم CNAS للسائق إن وُجد",
    "numero_casnos": "رقم CASNOS للسائق إن وُجد",
    "n_permis": "رقم رخصة السياقة",
    "date_delivrance_permis": "تاريخ إصدار رخصة السياقة YYYY-MM-DD",
    "categorie_permis": "B أو D أو غيرها",
    "n_certificat_formation": "رقم شهادة التكوين",
    "date_certificat_formation": "تاريخ شهادة التكوين YYYY-MM-DD",
    "n_telephone": "رقم الهاتف"
  },
  "vehicule": {
    "id_wilaya": 32,
    "activite": "Transports par Taxi",
    "nature_activite": "Individuelle أو inter-wilayas",
    "n_porte": "رقم الباب",
    "n_immatriculation": "رقم تسجيل المركبة (مثال: 05427.112.32)",
    "n_chasis": "الرقم التسلسلي VIN",
    "type": "طراز المركبة",
    "marque": "الصنف (RENAULT, TOYOTA إلخ)",
    "genre": "نوع المركبة",
    "carrossere": "الهيكل",
    "n_siege": "عدد المقاعد",
    "puissance": "القوة المحركة",
    "energie": "نوع الوقود (ESSENCE, GASOIL, GPL إلخ)",
    "annee_premier_utilisation": "سنة أول استعمال",
    "n_certificat_cerculation": "رقم رخصة السير",
    "date_fin_validite_certificat": "تاريخ انتهاء رخصة السير YYYY-MM-DD"
  },
  "licence_modjahid": {
    "id_wilaya": 32,
    "n_porte": "رقم الباب",
    "nin_modjahid": "رقم التعريف الوطني للمجاهد/الشهيد (إن وُجد في الوثائق)",
    "prenom": "اسم المجاهد/الشهيد بالفرنسية",
    "nom": "لقب المجاهد/الشهيد بالفرنسية",
    "prenom_ar": "اسم المجاهد/الشهيد بالعربية",
    "nom_ar": "لقب المجاهد/الشهيد بالعربية",
    "date_naissance": "تاريخ ميلاد المجاهد/الشهيد YYYY-MM-DD",
    "n_acte_naissance": "رقم شهادة ميلاد المجاهد",
    "type_licence": "أرملة مجاهد أو ابن شهيد أو ابنة شهيد إلخ",
    "n_decision_octroi": "رقم قرار المنح",
    "date_decision_octroi": "تاريخ قرار المنح YYYY-MM-DD"
  },
  "notes": "أي ملاحظات أو معلومات إضافية غير مصنفة"
}
"""


def check_dependencies():
    """التحقق من وجود المكتبات المطلوبة."""
    missing = []
    try:
        import anthropic
    except ImportError:
        missing.append("anthropic")
    try:
        import fitz
    except ImportError:
        missing.append("pymupdf")
    try:
        import openpyxl
    except ImportError:
        missing.append("openpyxl")

    if missing:
        print(f"❌ مكتبات مفقودة: {', '.join(missing)}")
        print(f"   شغّل: pip install {' '.join(missing)}")
        sys.exit(1)


def pdf_to_images_base64(pdf_path: str) -> list[dict]:
    """تحويل صفحات PDF إلى صور base64."""
    import fitz
    doc = fitz.open(pdf_path)
    images = []
    print(f"📄 عدد الصفحات: {len(doc)}")
    for i, page in enumerate(doc):
        mat = fitz.Matrix(2.0, 2.0)  # 2x zoom = ~144 DPI
        pix = page.get_pixmap(matrix=mat)
        img_data = pix.tobytes("jpeg")
        b64 = base64.standard_b64encode(img_data).decode("utf-8")
        images.append({
            "type": "image",
            "source": {
                "type": "base64",
                "media_type": "image/jpeg",
                "data": b64
            }
        })
        print(f"   ✓ صفحة {i+1}/{len(doc)}")
    doc.close()
    return images


def extract_data_with_claude(pdf_path: str) -> dict:
    """استخراج البيانات من PDF باستخدام Claude API."""
    import anthropic

    if not API_KEY:
        print("❌ مفتاح API غير موجود.")
        print("   أضفه في المتغير البيئي: set ANTHROPIC_API_KEY=sk-ant-...")
        print("   أو عدّل السطر API_KEY في أعلى الملف")
        sys.exit(1)

    client = anthropic.Anthropic(api_key=API_KEY)

    print("\n🔍 تحويل PDF إلى صور...")
    images = pdf_to_images_base64(pdf_path)

    print(f"\n🤖 إرسال إلى Claude ({MODEL}) للاستخراج...")
    content = images + [{"type": "text", "text": EXTRACTION_PROMPT}]

    message = client.messages.create(
        model=MODEL,
        max_tokens=4096,
        messages=[{"role": "user", "content": content}]
    )

    raw = message.content[0].text.strip()

    # استخراج JSON من الرد
    json_match = re.search(r'\{.*\}', raw, re.DOTALL)
    if json_match:
        raw = json_match.group(0)

    try:
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        print(f"❌ خطأ في تحليل JSON: {e}")
        print("--- رد النموذج ---")
        print(raw[:2000])
        sys.exit(1)

    return data


def get_column_map(ws) -> dict:
    """قراءة خريطة الأعمدة من ورقة Excel (الصف 1 = المفاتيح الفرنسية)."""
    col_map = {}
    for col_idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)), 1):
        if cell is not None:
            col_map[str(cell).strip()] = col_idx
    return col_map


def find_next_empty_row(ws) -> int:
    """إيجاد أول صف فارغ في ورقة العمل (البيانات تبدأ من الصف 3)."""
    for row_idx in range(3, ws.max_row + 2):
        if ws.cell(row=row_idx, column=1).value is None:
            return row_idx
    return ws.max_row + 1


def check_duplicate(ws, col_map: dict, key_field: str, key_value) -> bool:
    """فحص التكرار: هل هذا رقم الباب موجود مسبقاً؟"""
    if key_field not in col_map or key_value is None:
        return False
    col_idx = col_map[key_field]
    for row_idx in range(3, ws.max_row + 1):
        cell_val = ws.cell(row=row_idx, column=col_idx).value
        if cell_val is not None and str(cell_val).strip() == str(key_value).strip():
            return True
    return False


def write_row(ws, col_map: dict, data: dict, row_idx: int):
    """كتابة صف بيانات في ورقة العمل."""
    for field, value in data.items():
        if value is None:
            continue
        # البحث عن اسم العمود (مع تسامح بالمسافات والشرطات)
        col_idx = None
        for key in col_map:
            if key.replace("-", "_").replace(" ", "_").lower() == field.replace("-", "_").replace(" ", "_").lower():
                col_idx = col_map[key]
                break
        if col_idx is None:
            col_idx = col_map.get(field)
        if col_idx:
            ws.cell(row=row_idx, column=col_idx).value = value


def update_excel(excel_path: str, extracted: dict):
    """تحديث ملف Excel بالبيانات المستخرجة."""
    import openpyxl

    # نسخة احتياطية
    backup_path = excel_path.replace(".xlsx", f"_backup_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(excel_path, backup_path)
    print(f"\n💾 نسخة احتياطية: {Path(backup_path).name}")

    wb = openpyxl.load_workbook(excel_path)
    n_porte = str(extracted.get("operateur", {}).get("n-porte") or
                  extracted.get("vehicule", {}).get("n_porte") or
                  extracted.get("licence_modjahid", {}).get("n_porte") or "")

    sheets_config = [
        ("Operateur",        extracted.get("operateur", {}),        "n-porte",   n_porte),
        ("Chauffeur",        extracted.get("chauffeur", {}),         None,        None),
        ("Véhicule",         extracted.get("vehicule", {}),          "n_porte",   n_porte),
        ("Licence Modjahid", extracted.get("licence_modjahid", {}),  "n_porte",   n_porte),
    ]

    for sheet_name, data, dup_field, dup_val in sheets_config:
        if not data or sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        col_map = get_column_map(ws)

        if dup_field and check_duplicate(ws, col_map, dup_field, dup_val):
            print(f"⚠️  [{sheet_name}] رقم الباب {dup_val} موجود مسبقاً — تم التخطي")
            continue

        row_idx = find_next_empty_row(ws)
        write_row(ws, col_map, data, row_idx)
        print(f"✅ [{sheet_name}] تمت الإضافة في الصف {row_idx}")

    wb.save(excel_path)
    print(f"\n✅ تم حفظ الملف: {Path(excel_path).name}")


def print_preview(data: dict):
    """عرض ملخص البيانات المستخرجة."""
    print("\n" + "="*60)
    print("📋 ملخص البيانات المستخرجة:")
    print("="*60)

    op = data.get("operateur", {})
    ch = data.get("chauffeur", {})
    ve = data.get("vehicule", {})
    lm = data.get("licence_modjahid", {})

    print(f"\n🏠 صاحب/ة الرخصة (Operateur):")
    print(f"   الاسم: {op.get('prenom_ar','')} {op.get('nom_ar','')}")
    print(f"   NIN:   {op.get('nin','')}")
    print(f"   رقم الباب: {op.get('n-porte','')}")
    print(f"   الرخصة: من {op.get('date_debut_contrat','')} إلى {op.get('date_fin_contrat','')}")
    print(f"   الصفة: {op.get('qualite_oper_physique','')}")

    print(f"\n🚗 السائق (Chauffeur):")
    print(f"   الاسم: {ch.get('prenom_ar','')} {ch.get('nom_ar','')}")
    print(f"   الصفة: {ch.get('qualite_chauffeur','')}")
    print(f"   رخصة السياقة: {ch.get('n_permis','')}")
    print(f"   CNAS: {ch.get('numero_cnas','')}")

    print(f"\n🚙 المركبة (Véhicule):")
    print(f"   الصنف: {ve.get('marque','')} {ve.get('type','')}")
    print(f"   التسجيل: {ve.get('n_immatriculation','')}")
    print(f"   الرقم التسلسلي: {ve.get('n_chasis','')}")
    print(f"   الوقود: {ve.get('energie','')} | المقاعد: {ve.get('n_siege','')}")

    print(f"\n📜 رخصة المجاهد (Licence Modjahid):")
    print(f"   المجاهد/الشهيد: {lm.get('prenom_ar','')} {lm.get('nom_ar','')}")
    print(f"   نوع الرخصة: {lm.get('type_licence','')}")
    print(f"   قرار المنح: {lm.get('n_decision_octroi','')} بتاريخ {lm.get('date_decision_octroi','')}")

    if data.get("notes"):
        print(f"\n📝 ملاحظات: {data['notes']}")
    print("="*60)


def main():
    check_dependencies()

    if len(sys.argv) < 2:
        print(__doc__)
        print("مثال: python taxi_agent.py \"اسم_الملف.pdf\"")
        sys.exit(1)

    pdf_path = sys.argv[1]
    excel_path = sys.argv[2] if len(sys.argv) > 2 else EXCEL_PATH

    if not os.path.exists(pdf_path):
        print(f"❌ الملف غير موجود: {pdf_path}")
        sys.exit(1)

    if not os.path.exists(excel_path):
        print(f"❌ ملف Excel غير موجود: {excel_path}")
        print("   عدّل متغير EXCEL_PATH في أعلى الكود")
        sys.exit(1)

    print(f"📂 الملف: {pdf_path}")
    print(f"📊 Excel: {excel_path}")

    # الاستخراج
    data = extract_data_with_claude(pdf_path)

    # حفظ JSON للمراجعة
    json_out = pdf_path.replace(".pdf", "_extracted.json")
    with open(json_out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"\n💾 البيانات المستخرجة حُفظت في: {json_out}")

    # عرض الملخص
    print_preview(data)

    # تأكيد الإدخال
    print("\n❓ هل تريد إضافة هذه البيانات إلى ملف Excel؟ (y/n): ", end="")
    answer = input().strip().lower()
    if answer not in ("y", "yes", "نعم", "1"):
        print("❌ تم الإلغاء. لم يتم تعديل الملف.")
        sys.exit(0)

    # الكتابة في Excel
    update_excel(excel_path, data)

    print("\n🎉 تمت العملية بنجاح!")


if __name__ == "__main__":
    main()
