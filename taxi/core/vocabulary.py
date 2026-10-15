#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""تعلّم المصطلحات من ورقة Referances في ملف Excel نفسه.

مُستخرَج من taxi_pro.py v5.7.1 — المنطق كما هو، مُجرَّب على ملفات حقيقية.
"""
from __future__ import annotations

from taxi.constants import WILAYA_ID
from taxi.core.schema import (SCHEMA, REF_MAP, REF_HEADER_AS_VALUE,
                             ENUM_FIELDS, REF_SHEET)
from taxi.core.text import _norm
from taxi.excel.reference import read_reference_sheet


_VOCAB: dict = {}


def learn_vocabulary(excel_path: str, max_vals: int = 12) -> dict:
    """
    مصدران للمصطلحات، بهذا الترتيب:
      ① ورقة Referances — المرجعية الرسمية (لها الأولوية المطلقة)
      ② القيم المستعملة فعلاً في أوراق البيانات — لما لا تغطّيه المرجعية
    """
    import openpyxl
    from collections import Counter
    vocab: dict = {}
    valid = {(g, f) for g, grp in SCHEMA.items() for f, _, _ in grp["fields"]}

    # ① المرجعية الرسمية
    ref = read_reference_sheet(excel_path)
    for head, targets in REF_MAP.items():
        if head in ref:
            for g, f in targets:
                if (g, f) in valid:
                    vocab.setdefault(g, {})[f] = list(ref[head])
    for head, targets in REF_HEADER_AS_VALUE.items():
        if head in ref:
            for g, f in targets:
                if (g, f) in valid:
                    vocab.setdefault(g, {})[f] = [head]
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    except Exception:
        return vocab
    try:
        for gkey, grp in SCHEMA.items():
            sheet = grp["sheet"]
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 3:
                continue
            head = {str(c).strip(): i for i, c in enumerate(rows[0]) if c is not None}
            for fkey, _, _ in grp["fields"]:
                if (gkey, fkey) not in ENUM_FIELDS:
                    continue
                if vocab.get(gkey, {}).get(fkey):
                    continue          # المرجعية الرسمية تسبق الاستعمال
                ci = None
                for k, i in head.items():
                    if _norm(k) == _norm(fkey):
                        ci = i
                        break
                if ci is None:
                    continue
                cnt = Counter()
                for r in rows[2:]:
                    if ci < len(r) and r[ci] not in (None, ""):
                        cnt[str(r[ci]).strip()] += 1
                # نتجاهل القيم النادرة جداً (غالباً أخطاء إدخال سابقة)
                total = sum(cnt.values()) or 1
                keep = [v for v, c in cnt.most_common(max_vals) if c / total >= 0.004]
                if keep:
                    vocab.setdefault(gkey, {})[fkey] = keep
    finally:
        wb.close()
    return vocab
