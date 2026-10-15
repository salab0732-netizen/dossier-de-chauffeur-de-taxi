#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""الكتابة في Excel: مطابقة الأعمدة، كشف التكرار، نسخة احتياطية قبل كل كتابة.

مُستخرَج من taxi_pro.py v5.7.1 — المنطق كما هو، مُجرَّب على ملفات حقيقية.
"""
from __future__ import annotations

import datetime
import re
import shutil
from pathlib import Path

from taxi.core.schema import SCHEMA
from taxi.core.text import _norm
from taxi.core.enrich import section_active


def excel_column_map(ws) -> dict[str, int]:
    row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    out = {}
    for ci, cell in enumerate(row1, 1):
        if cell is not None:
            out[str(cell).strip()] = ci
    return out



def find_column(col_map: dict[str, int], field: str) -> int | None:
    if field in col_map:
        return col_map[field]
    target = _norm(field)
    for k, ci in col_map.items():
        if _norm(k) == target:
            return ci
    return None


def next_empty_row(ws, ncols: int = 0) -> int:
    """
    أول صف فارغ تماماً.
    ⚠️ لا يكفي فحص العمود الأول: سجلّات الشركات تترك id_wilaya فارغاً
    بينما تملأ عشرين عموداً أخرى — فحصُ عمود واحد يدمّرها.
    """
    ncols = ncols or ws.max_column
    for ri in range(3, ws.max_row + 2):
        if all(ws.cell(row=ri, column=ci).value in (None, "")
               for ci in range(1, ncols + 1)):
            return ri
    return ws.max_row + 1


def duplicate_row(ws, col_map, chains, data) -> tuple[int | None, str]:
    """
    يبحث عن صف مطابق باستعمال سلاسل مفاتيح مرتّبة.
    تُجرَّب كل سلسلة بالترتيب؛ تُتخطّى أي سلسلة ينقص أحد حقولها.
    يعيد (رقم الصف أو None، وصف المفتاح المستعمل).
    """
    if isinstance(chains, str):
        chains = [(chains,)]
    for chain in chains or []:
        cols, vals = [], []
        for f in chain:
            ci = find_column(col_map, f)
            v = data.get(f)
            if not ci or v in (None, "", "null"):
                cols = []
                break
            cols.append(ci)
            vals.append(str(v).strip())
        if not cols:
            continue
        for ri in range(3, ws.max_row + 1):
            if all(ws.cell(row=ri, column=c).value is not None
                   and str(ws.cell(row=ri, column=c).value).strip() == t
                   for c, t in zip(cols, vals)):
                return ri, "+".join(chain)
    return None, ""


def write_to_excel(excel_path: str, data: dict, overwrite_dups: bool, log) -> dict:
    """كتابة البيانات في Excel. يعيد ملخّص العملية."""
    import openpyxl

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = str(Path(excel_path).with_name(f"{Path(excel_path).stem}_backup_{ts}.xlsx"))
    shutil.copy2(excel_path, backup)
    log(f"💾 نسخة احتياطية: {Path(backup).name}")

    wb = openpyxl.load_workbook(excel_path)
    summary = {"added": [], "updated": [], "skipped": [], "missing": [], "backup": backup}

    for gkey, grp in SCHEMA.items():
        row_data = {k: v for k, v in (data.get(gkey) or {}).items()
                    if v not in (None, "", "null")}
        if not row_data:
            continue
        # قسم لا يحمل إلا ثوابت مشتقّة ليس سجلاً — لا يُكتب
        if not section_active(data, gkey):
            log(f"⏭️  [{grp['sheet']}] لا معلومة جوهرية — لم يُكتب")
            continue
        sheet = grp["sheet"]
        if sheet not in wb.sheetnames:
            summary["missing"].append(sheet)
            log(f"⚠️  الورقة «{sheet}» غير موجودة — تم التخطي")
            continue

        ws = wb[sheet]
        col_map = excel_column_map(ws)

        dup_ri, by = duplicate_row(ws, col_map, grp.get("dup") or grp.get("key"), row_data)

        if dup_ri and not overwrite_dups:
            summary["skipped"].append((sheet, dup_ri))
            log(f"⏭️  [{sheet}] سجل مطابق في الصف {dup_ri} (بـ {by}) — تم التخطي")
            continue
        if not dup_ri and not any(
                find_column(col_map, f) and row_data.get(f) not in (None, "", "null")
                for chain in (grp.get("dup") or []) for f in chain):
            summary.setdefault("nokey", []).append(sheet)
            log(f"⚠️  [{sheet}] لا مفتاح صالح لكشف التكرار — قد يتكرّر السجل")

        ri = dup_ri or next_empty_row(ws)
        written = 0
        for field, value in row_data.items():
            ci = find_column(col_map, field)
            if ci:
                ws.cell(row=ri, column=ci).value = value
                written += 1
        if dup_ri:
            summary["updated"].append((sheet, ri, written))
            log(f"♻️  [{sheet}] تحديث الصف {ri} — {written} حقل")
        else:
            summary["added"].append((sheet, ri, written))
            log(f"✅ [{sheet}] إضافة في الصف {ri} — {written} حقل")

    wb.save(excel_path)
    log(f"💾 تم حفظ: {Path(excel_path).name}")
    return summary
