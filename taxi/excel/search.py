#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""البحث في السجلّات المُدخَلة وإحصاؤها.

وحدة جديدة في v7: الواجهة القديمة كانت تكتب في Excel ولا تقرأ منه،
فلم يكن للموظّف سبيل لمعرفة ما أُدخل سابقاً إلا بفتح الملف بنفسه.
"""
from __future__ import annotations

from pathlib import Path

from taxi.core.schema import SCHEMA
from taxi.core.text import _norm

# الحقول التي يبحث بها الموظّف فعلاً، مرتّبةً بأولوية العرض
SEARCH_FIELDS = {
    "operateur": ["n-porte", "nom_ar", "prenom_ar", "nin", "numero_telephone"],
    "chauffeur": ["nin", "nom_ar", "prenom_ar", "n_permis"],
    "vehicule":  ["n_porte", "n_immatriculation", "marque", "n_chasis"],
    "licence_modjahid": ["n_porte", "nin_modjahid", "nom_ar", "prenom_ar"],
}


def _open(excel_path: str):
    import openpyxl
    return openpyxl.load_workbook(excel_path, data_only=True, read_only=True)


def _all_labels() -> set[str]:
    """كل التسميات العربية في المخطّط — للتعرّف على صفّ التسميات."""
    return {_norm(f[1]) for g in SCHEMA.values() for f in g["fields"]}


_LABELS = _all_labels()


def _is_label_row(row) -> bool:
    """هل هذا صفّ تسميات لا بيانات؟

    الملف الرسمي يحمل رأسَين: الصفّ 1 مفاتيح الأعمدة بالحروف اللاتينية،
    والصفّ 2 التسميات العربية. فتبدأ البيانات من الصفّ 3، وعدّ الصفّ 2
    سجلّاً يُظهر «رقم الباب | اللقب (AR)» كأنه متعامل.
    """
    vals = [_norm(v) for v in row if v not in (None, "")]
    if len(vals) < 3:
        return False
    # ثلاث خانات تُطابق تسميةً عربية بالحرف = صفّ تسميات. السجلّ الحقيقي
    # لا تُطابق فيه ثلاث قيم «رقم الولاية» و«النشاط» و«الصفة» بالحرف.
    # ولا تصلح نسبةٌ من عدد الخانات: الملف الرسمي فيه أعمدة خارج المخطّط،
    # فتهبط النسبة دون النصف مع أن الصفّ تسميات (19 من 40 في ورقة السائق).
    return sum(v in _LABELS for v in vals) >= 3


def _sheet_rows(ws, fields: list[str]) -> tuple[list[str], list[list]]:
    """يعيد (أسماء الأعمدة، الصفوف غير الفارغة) للحقول المطلوبة فقط."""
    rows = ws.iter_rows(values_only=True)
    try:
        head = next(rows)
    except StopIteration:
        return [], []

    # موقع كل حقل مطلوب في رأس الورقة
    pos: dict[str, int] = {}
    for i, h in enumerate(head):
        n = _norm(h)
        if not n:
            continue
        for f in fields:
            if f not in pos and _norm(f) == n:
                pos[f] = i

    cols = [f for f in fields if f in pos]
    out = []
    for r in rows:
        if _is_label_row(r):
            continue
        vals = [r[pos[f]] if pos[f] < len(r) else None for f in cols]
        if any(v not in (None, "") for v in vals):
            out.append(["" if v is None else str(v).strip() for v in vals])
    return cols, out


def search(excel_path: str, query: str = "", group: str = "operateur",
           limit: int = 200) -> dict:
    """بحث نصّي حرّ في ورقة واحدة.

    يُطابَق بـ_norm فلا يُفرّق الفراغ ولا حالة الحرف، ويُطابَق جزئياً
    لأن الموظّف يكتب «2253» أو «بلفضال» لا القيمة كاملةً.
    """
    grp = SCHEMA.get(group)
    if not grp:
        return {"error": f"قسم غير معروف: {group}", "rows": []}

    fields = SEARCH_FIELDS.get(group) or [f[0] for f in grp["fields"]][:6]
    try:
        wb = _open(excel_path)
    except FileNotFoundError:
        return {"error": "ملف Excel غير موجود", "rows": []}
    except Exception as e:
        return {"error": f"تعذّر فتح الملف: {e}", "rows": []}

    try:
        if grp["sheet"] not in wb.sheetnames:
            return {"error": f"الورقة «{grp['sheet']}» غير موجودة", "rows": []}
        cols, rows = _sheet_rows(wb[grp["sheet"]], fields)
    finally:
        wb.close()

    q = _norm(query)
    if q:
        rows = [r for r in rows if any(q in _norm(v) for v in r)]

    labels = {f[0]: f[1] for f in grp["fields"]}
    return {
        "sheet":   grp["sheet"],
        "title":   grp["title"],
        "columns": [{"key": c, "label": labels.get(c, c)} for c in cols],
        "rows":    rows[:limit],
        "total":   len(rows),
        "shown":   min(len(rows), limit),
    }


def stats(excel_path: str) -> dict:
    """عدّ السجلّات في كل ورقة — للوحة المتابعة."""
    try:
        wb = _open(excel_path)
    except Exception as e:
        return {"error": str(e), "sheets": []}

    out = []
    try:
        seen = set()
        for key, grp in SCHEMA.items():
            sh = grp["sheet"]
            if sh in seen or sh not in wb.sheetnames:
                continue
            seen.add(sh)
            fields = SEARCH_FIELDS.get(key) or [f[0] for f in grp["fields"]][:4]
            _, rows = _sheet_rows(wb[sh], fields)
            out.append({"sheet": sh, "title": grp["title"],
                        "icon": grp.get("icon", ""), "count": len(rows)})
    finally:
        wb.close()

    return {"sheets": out, "file": Path(excel_path).name}
