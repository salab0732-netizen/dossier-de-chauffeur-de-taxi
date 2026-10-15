#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""قراءة ورقة Referances — المرجعية الرسمية للمصطلحات.

مُستخرَج من taxi_pro.py v5.7.1 — المنطق كما هو، مُجرَّب على ملفات حقيقية.
"""
from __future__ import annotations

from taxi.core.schema import REF_SHEET


def read_reference_sheet(excel_path: str) -> dict[str, list[str]]:
    """يقرأ ورقة Referances: {عنوان العمود: [القيم المسموحة]}."""
    import openpyxl
    out: dict[str, list[str]] = {}
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    except Exception:
        return out
    try:
        if REF_SHEET not in wb.sheetnames:
            return out
        rows = list(wb[REF_SHEET].iter_rows(values_only=True))
        if len(rows) < 2:
            return out
        ncols = max(len(r) for r in rows)
        for ci in range(ncols):
            head = rows[0][ci] if ci < len(rows[0]) else None
            if head is None or str(head).strip() in ("", "_"):
                continue
            vals: list[str] = []
            # الصف 2 قد يحمل أول قيمة مسموحة
            v2 = rows[1][ci] if ci < len(rows[1]) else None
            if v2 is not None and str(v2).strip() not in ("", "_"):
                vals.append(str(v2).strip())
            for r in rows[2:]:
                v = r[ci] if ci < len(r) else None
                if v is not None and str(v).strip() not in ("", "_"):
                    s = str(v).strip()
                    if s not in vals:
                        vals.append(s)
            if vals:
                out[str(head).strip()] = vals
    finally:
        wb.close()
    return out
