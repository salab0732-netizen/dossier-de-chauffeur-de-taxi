#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
════════════════════════════════════════════════════════════════
  استخراج PDF محلي بدون إنترنت — Tesseract OCR
  Offline PDF Text Extractor using Tesseract
  ولاية البيض (32) — سيارات الأجرة
════════════════════════════════════════════════════════════════
  الاستخدام:
    python offline_ocr.py                  # معالجة كل PDF في المجلد
    python offline_ocr.py ملف.pdf          # معالجة ملف واحد
    python offline_ocr.py --setup          # تثبيت اللغة العربية

  المتطلبات:
    pip install pymupdf openpyxl
    + Tesseract مثبّت: https://github.com/UB-Mannheim/tesseract/wiki
════════════════════════════════════════════════════════════════
"""
from __future__ import annotations

import argparse, json, os, re, shutil, subprocess, sys, tempfile, urllib.request
from pathlib import Path
from datetime import datetime

# ──────────────────────────────────────────────────────────────────
SCRIPT_DIR   = Path(__file__).resolve().parent
TESSDATA_DIR = SCRIPT_DIR.parent / "taxi_tessdata"   # مسار ASCII — Tesseract لا يقبل مسارات عربية
EXCEL_PATH   = SCRIPT_DIR / "TAXI 32000.xlsx"
DPI          = 144          # نفس دقة البرنامج الرئيسي (zoom 2.0) — اختُبرت عليها القواعد
LANG         = None         # None = تلقائي: ara+fra إن وُجد fra.traineddata، وإلا ara
WILAYA       = 32

TESSDATA_URLS = {
    "ara": "https://github.com/tesseract-ocr/tessdata/raw/main/ara.traineddata",
    "fra": "https://github.com/tesseract-ocr/tessdata_fast/raw/main/fra.traineddata",
}

# ──────────────────────────────────────────────────────────────────
#  البحث عن Tesseract
# ──────────────────────────────────────────────────────────────────

def find_tesseract() -> str:
    found = shutil.which("tesseract")
    if found:
        return found
    for c in [r"C:\Program Files\Tesseract-OCR\tesseract.exe",
              r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"]:
        if Path(c).exists():
            return c
    return "tesseract"

def find_tessdata() -> str | None:
    for d in [TESSDATA_DIR, SCRIPT_DIR / "tessdata",
              Path(r"C:\Program Files\Tesseract-OCR\tessdata"),
              Path(r"C:\Program Files (x86)\Tesseract-OCR\tessdata")]:
        if (Path(d) / "ara.traineddata").exists():
            return str(d)
    return None

def pick_lang(tessdata: str | None) -> str:
    if LANG:
        return LANG
    if tessdata and (Path(tessdata) / "fra.traineddata").exists():
        return "ara+fra"
    return "ara"

# ──────────────────────────────────────────────────────────────────
#  تثبيت اللغة
# ──────────────────────────────────────────────────────────────────

def setup_arabic():
    """تنزيل ara و fra (الفرنسية للأسماء اللاتينية وبيانات البطاقة الرمادية)."""
    dest_dir = Path(find_tessdata() or TESSDATA_DIR)
    dest_dir.mkdir(parents=True, exist_ok=True)
    for lng, url in TESSDATA_URLS.items():
        dest = dest_dir / f"{lng}.traineddata"
        if dest.exists():
            print(f"✅ {lng} مثبّتة بالفعل: {dest}")
            continue
        print(f"⬇  تنزيل {lng}.traineddata ({url})")
        try:
            urllib.request.urlretrieve(url, str(dest))
            print(f"✅ تمّ التثبيت: {dest}")
        except Exception as e:
            print(f"❌ فشل التنزيل: {e}")
            sys.exit(1)

# ──────────────────────────────────────────────────────────────────
#  OCR صفحة واحدة
# ──────────────────────────────────────────────────────────────────

def ocr_image(img_bytes: bytes, tess_exe: str, tessdata: str | None, lang: str,
              with_tsv: bool = False):
    """نص الصفحة؛ ومع with_tsv: (النص، مواقع الكلمات TSV) من تمريرة واحدة."""
    with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as fh:
        fh.write(img_bytes)
        tmp_in = fh.name
    tmp_out = tmp_in[:-4]
    cmd = [tess_exe, tmp_in, tmp_out, "-l", lang, "--psm", "6",
           "-c", "tessedit_create_txt=1", "-c", f"tessedit_create_tsv={int(with_tsv)}"]
    if tessdata:
        cmd += ["--tessdata-dir", tessdata]
    try:
        subprocess.run(cmd, capture_output=True, timeout=120)
    except subprocess.TimeoutExpired:
        print("⏱", end="", flush=True)    # صفحة ضجيج طويلة — تُتجاوز بدل إسقاط الملف
    res = []
    for ext in (".txt", ".tsv"):
        f = tmp_out + ext
        t = ""
        if os.path.exists(f):
            with open(f, "r", encoding="utf-8") as fh:
                t = fh.read()
            os.remove(f)
        res.append(t)
    try:
        os.remove(tmp_in)
    except OSError:
        pass
    return (res[0], res[1]) if with_tsv else res[0]

# ──────────────────────────────────────────────────────────────────
#  استخراج الحقول بالـ Regex
# ──────────────────────────────────────────────────────────────────

def rx(text: str, *patterns, default=None):
    for p in patterns:
        m = re.search(p, text, re.MULTILINE)
        if m:
            return m.group(1).strip()
    return default

def extract_fields(text: str) -> dict:
    """واجهة قديمة — تُحيل إلى extract_taxi (الصفحات مفصولة بـ \\f)."""
    return extract_taxi(text, WILAYA)


# ══════════════════════════════════════════════════════════════════════
#  استخراج الحقول من نص OCR — يعمل صفحةً صفحة حسب نوع الوثيقة
#  (رخصة الاستغلال · البطاقة الرمادية · بطاقة التعريف · CNAS · CASNOS ·
#   شهادة التكوين · قرار المجاهدين · رخصة السياقة)
#  الصفحات مفصولة بالمحرف \f.  كل قيمة لا تُقرأ بثقة تُترك فارغة.
# ══════════════════════════════════════════════════════════════════════
import re as _re
import datetime as _dt
import difflib as _difflib

_BIDI = dict.fromkeys(map(ord, "‎‏‪‫‬‭‮"
                                "⁦⁧⁨⁩﻿​"), None)
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")

_COMMUNES_32 = [
    "البيض", "الأبيض سيد الشيخ", "بوعلام", "الرقاصة", "بريزينة", "الغاسول",
    "استيتن", "الشقيق", "بوقطب", "الكاف الأحمر", "عين العراك", "اربوات",
    "الشلالة", "بوسمغون", "المهرة", "الخيثر", "توسمولين", "سيدي سليمان",
    "الكراكدة", "سيدي عمر", "سيدي الطيفور", "سيدي عامر", "الأبيض سيدي الشيخ",
    "الكاف لحمر", "عين العراك",
]

# «تاريخ الميلاد» كما يشوّهه OCR أحياناً
_DOB = r"(?:الميلاد|المبلاد|اليلاد|المهلاد|الميلد|الملاد|شيلاد)"

_GROUPAGE = {"A+": "A +", "A-": "A -", "B+": "B + ", "B-": "B -",
             "AB+": "AB + ", "AB-": "AB -", "O+": "O +", "O-": "O -"}

_MARQUES = ["RENAULT", "DACIA", "HYUNDAI", "PEUGEOT", "CHEVROLET", "TOYOTA",
            "KIA", "SUZUKI", "VOLKSWAGEN", "CITROEN", "FIAT", "SEAT", "OPEL",
            "FORD", "NISSAN", "SKODA", "MAZDA", "MITSUBISHI", "GEELY", "CHERY",
            "DAEWOO", "HONDA", "ISUZU", "MERCEDES", "BMW", "AUDI", "LADA",
            "SSANGYONG", "JAC", "BYD", "DFSK", "BAIC", "DONGFENG"]
_MARQUES_AR = {"رونو": "RENAULT", "داسيا": "DACIA", "هيونداي": "HYUNDAI",
               "بيجو": "PEUGEOT", "شفروليه": "CHEVROLET", "شيفروليه": "CHEVROLET",
               "تويوتا": "TOYOTA", "كيا": "KIA", "سوزوكي": "SUZUKI"}


def _clean(t: str) -> str:
    return (t or "").translate(_BIDI).translate(_DIGITS)


def _ar(s: str) -> str:
    """تطبيع عربي للمقارنة فقط."""
    s = _re.sub(r"[ًٌٍَُِّْـ]", "", s or "")
    s = _re.sub(r"[أإآٱ]", "ا", s)
    s = s.replace("ة", "ه").replace("ى", "ي").replace("ؤ", "و").replace("ئ", "ي")
    return _re.sub(r"[^ء-ي]", "", s)


def _similar(a: str, b: str) -> float:
    a, b = _ar(a), _ar(b)
    if not a or not b:
        return 0.0
    return _difflib.SequenceMatcher(None, a, b).ratio()


def _is_ar_word(w: str) -> bool:
    return bool(w) and bool(_re.fullmatch(r"[ء-ي]{2,}", w))


def _ar_words(s: str, maxn: int = 3) -> str | None:
    """أول كلمات عربية متتالية (حتى maxn) — يتوقف عند أول رمز غريب."""
    out = []
    for w in _re.split(r"[\s:.,؛،]+", (s or "").strip()):
        if not w:
            continue
        if _is_ar_word(w):
            out.append(w)
            if len(out) >= maxn:
                break
        else:
            break
    return " ".join(out) or None


def _date(y, m, d, lo=1900, hi=2040):
    try:
        y, m, d = int(y), int(m), int(d)
        if lo <= y <= hi:
            return _dt.date(y, m, d).isoformat()
    except (ValueError, TypeError):
        pass
    return None


def _dates(s: str, lo=1900, hi=2040) -> list[str]:
    """كل التواريخ الصالحة في النص بالترتيب (YYYY.MM.DD أو DD.MM.YYYY)."""
    out = []
    for m in _re.finditer(r"(?<!\d)(\d{4})\s?[./\-]\s?(\d{1,2})\s?[./\-]\s?(\d{1,2})(?!\d)"
                          r"|(?<!\d)(\d{1,2})\s?[./\-]\s?(\d{1,2})\s?[./\-]\s?(\d{4})(?!\d)", s):
        if m.group(1):
            v = _date(m.group(1), m.group(2), m.group(3), lo, hi)
        else:
            v = _date(m.group(6), m.group(5), m.group(4), lo, hi)
        if v:
            out.append(v)
    return out


def _title(s: str | None) -> str | None:
    if not s:
        return None
    s = _re.sub(r"[^A-Za-z\-' ]", " ", s)
    s = " ".join(w for w in s.split() if len(w) >= 2)
    if not s or len(s) < 2:
        return None
    return " ".join(p.capitalize() for p in s.split())


def _latin_after(label_rx: str, page: str, maxw: int = 3) -> str | None:
    m = _re.search(label_rx + r"\s*[:.]?\s*((?:[A-Z][A-Za-z\-']+\s*){1,%d})" % maxw, page)
    if not m:
        return None
    stop = {"FILS", "FILLE", "ET", "DE", "NOM", "PRENOM", "DATE", "LIEU", "N", "NE", "NEE"}
    words = []
    for w in m.group(1).split():
        if w.upper() in stop or not _re.fullmatch(r"[A-Z][A-Z\-']+|[A-Z][a-z\-']+", w):
            break
        # بعد كلمة بحروف كبيرة لا نقبل إلا كلمات بحروف كبيرة (≥3) — الباقي ضجيج
        if words and words[0].isupper() and not (w.isupper() and len(w) >= 3):
            break
        words.append(w)
    return _title(" ".join(words[:maxw])) if words else None


def _nins(s: str) -> list[str]:
    out = []
    for line in s.split("\n"):
        j = _re.sub(r"[\s.]+", "", line)
        for m in _re.finditer(r"(?<!\d)([12][01]\d{16})(?!\d)", j):
            if m.group(1) not in out:
                out.append(m.group(1))
    return out


def _nins_loose(s: str) -> list[str]:
    """نوافذ 18 خانة داخل سلاسل من 19-20 رقماً (رقم زائد ألصقه OCR).
    لا تُستعمل إلا إذا وافقت سنة الميلاد المعروفة."""
    out = []
    for line in s.split("\n"):
        j = _re.sub(r"[\s.]+", "", line)
        for m in _re.finditer(r"(?<!\d)(\d{19,20})(?!\d)", j):
            run = m.group(1)
            for i in range(len(run) - 17):
                c = run[i:i + 18]
                if _re.match(r"[12][01]", c) and c not in out:
                    out.append(c)
    return out


def _nin_year(nin: str) -> int | None:
    try:
        return 1000 + int(nin[2:5])
    except (ValueError, TypeError):
        return None


def _commune(s: str) -> str | None:
    s = _ar_words(s, 4)
    if not s:
        return None
    best, score = None, 0.0
    for c in _COMMUNES_32:
        for k in range(1, 4):
            cand = " ".join(s.split()[:k])
            r = _similar(cand, c)
            if r > score:
                best, score = c, r
    return best if score >= 0.8 else None


def _parents(s: str, nom: str = ""):
    """«سليمان بن بوبكر و شريطات عائشة بنت ...» ← (الأب، لقب الأم، اسم الأم)."""
    s = _re.sub(r"[^ء-ي\s]", " ", s or "")
    parts = _re.split(r"\s+و\s*", s, maxsplit=1)
    pere = None
    if parts and parts[0].strip():
        w = parts[0].split()
        if len(w) > 1 and nom and _similar(w[0], nom) >= 0.8:
            w = w[1:]          # «راجع أحمد» ← اللقب ثم اسم الأب
        if w and w[0] in ("عبد", "ابو", "أبو") and len(w) > 1:
            pere = w[0] + " " + w[1]
        elif w and _is_ar_word(w[0]) and w[0] not in ("بن", "ابن"):
            pere = w[0]
    nom_m = pre_m = None
    if len(parts) > 1:
        w = parts[1].split()
        w = w[:w.index("بنت")] if "بنت" in w else w
        if w and w[0] in ("بن", "ابن") and len(w) > 2:
            w = [w[0] + " " + w[1]] + w[2:]
        if len(w) >= 2 and all(_is_ar_word(x.replace(" ", "")) for x in w[:2]):
            nom_m, pre_m = w[0], w[1]
    return pere, nom_m, pre_m


# ── تصنيف الصفحات ────────────────────────────────────────────────────
def _kind(p: str) -> set[str]:
    k = set()
    if _re.search(r"رخصة\s*مؤقتة|رقم\s*الباب", p):
        k.add("lic")
    if _re.search(r"التسلسلي|الحمولة|CARROSSERIE|GENRE\s+MARQUE|IMMATRICULATION|"
                  r"QUITTANCE|الرقم\s*السا|سنة\s*[اأ]ول\s*است", p):
        k.add("cg")
    if _re.search(r"بطا\S*\s*التعريف|سلطة\s*الإصدار|لطة\s*الإصدار", p):
        k.add("id")
    if _re.search(r"لغير\s*الأجراء|لغير\s*الاجراء|CASNOS|NON\s+SALARI", p):
        k.add("casnos")
    elif _re.search(r"للتأمينات|للتامينات|CNAS|الأجراء", p):
        k.add("cnas")
    if _re.search(r"دفتر\s*\S*\s*مقاعد|دقتر\s*مقاعد|الدورة\s*رقم", p):
        k.add("form")
    if _re.search(r"بصفت", p) and _re.search(r"مجاهد|شهيد|مكان\s*الاست", p):
        k.add("dec")
    if _re.search(r"رخصة\s*ال\S*ياقة|PERMIS\s+DE\s+CONDUIRE|DRIVING", p) and "lic" not in k:
        k.add("permis")
    return k


def _brand(s: str, short_ok: bool) -> str | None:
    """الصنف من نص (بحروف كبيرة). الأسماء القصيرة (KIA…) لا تُقبل إلا في
    سطر الصنف نفسه (short_ok) لأنها تظهر صدفةً في ضجيج OCR."""
    for mk in _MARQUES:
        if (short_ok or len(mk) >= 5) and _re.search(r"\b" + mk + r"\b", s):
            return mk
    # مطابقة متسامحة مع أخطاء OCR («RENANLT»، «CHEVR OLETE»)
    ws = _re.findall(r"\b[A-Z]{3,10}\b", s)
    best, score = None, 0.0
    for i, w in enumerate(ws):
        for cand in {w, w + ws[i + 1] if i + 1 < len(ws) else w}:
            if len(cand) < 5:
                continue
            for mk in _MARQUES:
                if len(mk) < 5:
                    continue
                r = _difflib.SequenceMatcher(None, cand, mk).ratio()
                if r > score:
                    best, score = mk, r
    return best if score >= 0.8 else None


def _carrosserie(s: str) -> str | None:
    """الهيكل كما هو مطبوع في البطاقة (BREAK / BERLINE / C,I) — توحيده مع
    قيم ملف Excel يتولاه فحص البرنامج بعد الاستخراج."""
    s = (s or "").upper()
    if _re.search(r"\bB?R[EI]?AK\b|BREAK", s):
        return "BREAK"
    if _re.search(r"BERLI|برلين", s):
        return "BERLINE"
    if _re.search(r"\bC\.?\s?I\b|ق\.?\s?د", s):
        return "C,I"
    return None


def _tsv_words(tsv: str) -> list[dict]:
    out = []
    for line in (tsv or "").split("\n"):
        c = line.split("\t")
        if len(c) < 12 or c[0] != "5":
            continue
        t = _clean(c[11]).strip()
        if not t:
            continue
        try:
            l, tp, w, h = int(c[6]), int(c[7]), int(c[8]), int(c[9])
        except ValueError:
            continue
        out.append({"l": l, "t": tp, "r": l + w, "b": tp + h, "h": h, "text": t,
                    "U": _re.sub(r"[^A-Z0-9]", "", t.upper()
                                 .replace("É", "E").replace("È", "E"))})
    return out


_CG_LABELS = {          # كلمة العنوان المطبوعة تحت القيمة ← الحقل
    "MARQUE": "marque", "PUISSANCE": "puissance", "ENERGIE": "energie",
    "PLACES": "n_siege", "ASSISES": "n_siege", "CARROSSERIE": "carrossere",
    "SERIE": "n_chasis", "ANNEE": "annee_premier_utilisation",
    "PREMIERE": "annee_premier_utilisation", "UTILE": "charge utile",
    "POIDS": "poids total", "TYPE": "type",
}
_LABEL_WORDS = set(_CG_LABELS) | {"GENRE", "DANS", "LA", "DU", "N", "TOTAL", "EN",
                                  "CHARGE", "DE", "MISE", "CIRCULATION", "NUMERO",
                                  "PRECEDENT", "DIMMATRICULATION", "IMMATRICULATION"}


def _cg_cells(tsv: str) -> dict:
    """قراءة جدول البطاقة الرمادية بالمواقع: لكل عنوان مطبوع (MARQUE، PUISSANCE…)
    نأخذ الكلمات التي فوقه مباشرة في عموده. كل كلمة تُنسب إلى أقرب عنوان في
    الصف نفسه، فلا تختلط قيمة خانة بجارتها."""
    ws = _tsv_words(tsv)
    if not ws:
        return {}

    # ① مراسي العناوين: (الحقل، الحدّ الأيسر، الأيمن، الأعلى، الارتفاع)
    anchors = []
    used = set()
    for lab in ws:
        U = lab["U"]
        if id(lab) in used:
            continue
        line = [w for w in ws if abs(w["t"] - lab["t"]) <= 15]
        f = None
        group = [lab]
        if U == "SERIE":
            f = "n_chasis"
            group = [w for w in line if (w["U"] in ("N", "DANS", "LA", "SERIE", "DU")
                                         and abs(w["l"] - lab["l"]) < 260)
                     or (w["U"] == "TYPE" and 0 < w["l"] - lab["l"] < 200)]
        elif U == "TYPE":
            ser = [w for w in line if w["U"] == "SERIE"]
            if ser and 0 < lab["l"] - ser[0]["l"] < 200:
                continue                         # جزء من «SERIE DU TYPE»
            f = "type"
        elif U in ("ANNEE", "PREMIERE"):
            f = "annee_premier_utilisation"
            group = [w for w in line if w["U"] in ("ANNEE", "DE", "PREMIERE")
                     and abs(w["l"] - lab["l"]) < 220]
        elif U in ("PLACES", "ASSISES"):
            f = "n_siege"
        elif U in ("MARQUE", "PUISSANCE", "ENERGIE", "CARROSSERIE"):
            f = _CG_LABELS[U]
        elif U in ("GENRE", "POIDS", "UTILE", "PRECEDENT", "DIMMATRICULATION",
                   "IMMATRICULATION"):
            f = "_" + U                          # مرساة بلا حقل: تحدّ الجيران فقط
        if not f:
            continue
        for w in group:
            used.add(id(w))
        if any(an["f"] == f for an in anchors):
            continue
        anchors.append({"f": f, "l": min(w["l"] for w in group),
                        "r": max(w["r"] for w in group), "t": min(w["t"] for w in group),
                        "h": max(w["h"] for w in group)})
    for an in anchors:
        an["cx"] = (an["l"] + an["r"]) / 2

    labelish = _LABEL_WORDS | {"PLACE", "ASSISE"}

    # ارتفاع سطر العناوين المطبوعة (≈17 نقطة عند تكبير 2) — مقياس الصفحة
    hs = sorted(an["h"] for an in anchors) or [18]
    hh = max(12, hs[len(hs) // 2])

    def value_of(an: dict) -> str | None:
        row = [o for o in anchors if abs(o["t"] - an["t"]) <= 2.5 * hh]
        lo = an["t"] - 5.5 * hh
        cands = []
        for w in ws:
            cy = (w["t"] + w["b"]) / 2       # القيمة بخط عريض قد تلامس العنوان
            if w["U"] in labelish or not (lo <= cy <= an["t"] + 0.5 * hh):
                continue
            cx = (w["l"] + w["r"]) / 2
            if not (an["l"] - 45 <= cx <= an["r"] + 45):
                continue
            nearest = min(row, key=lambda o: 0 if o["l"] <= cx <= o["r"] else
                          min(abs(cx - o["l"]), abs(cx - o["r"])))
            if nearest is an:
                cands.append(w)
        if not cands:
            return None
        cyc = lambda c: (c["t"] + c["b"]) / 2
        near = max(cyc(c) for c in cands)      # الأقرب إلى العنوان
        line = sorted([c for c in cands if near - cyc(c) <= max(20, 1.2 * hh)],
                      key=lambda c: c["l"])
        return " ".join(c["text"] for c in line)

    out: dict = {}
    for an in anchors:
        f = an["f"]
        if f.startswith("_"):
            continue
        v = value_of(an)
        if not v:
            continue
        VU = v.upper()
        V = _re.sub(r"[^A-Z0-9]", "", VU)
        if f == "marque":
            mk = _brand(VU, short_ok=True)
            if mk:
                out[f] = mk
        elif f == "type":
            if 5 <= len(V) <= 12 and _re.search(r"\d", V) and _re.search(r"[A-Z]", V):
                out[f] = V
        elif f == "n_chasis":
            if len(V) == 17:
                out[f] = V.translate(str.maketrans("IOQ", "100"))   # لا I/O/Q في VIN
            elif len(V) == 8 and len(_re.sub(r"\D", "", V)) >= 5:
                out[f] = V
        elif f == "puissance":
            d = _re.sub(r"\D", "", VU)
            if _re.fullmatch(r"0?\d{1,2}", d or "x") and 3 <= int(d) <= 20:
                out[f] = f"{int(d)} CHEVEAUX"
        elif f == "n_siege":
            d = _re.sub(r"\D", "", VU)
            if _re.fullmatch(r"0?[2-9]", d or "x"):
                out[f] = str(int(d))
        elif f == "energie":
            if _re.search(r"\bGO\b|GASOIL|مازوت", VU):
                out[f] = "Diesel"
            elif _re.search(r"GPL|غاز|\bGP\b", VU):
                out[f] = "GPL"
            elif _re.search(r"\bES\b|\bEG\b|ESS|بنزين", VU):
                out[f] = "Essence"
        elif f == "carrossere":
            c = _carrosserie(VU)
            if c:
                out[f] = c
        elif f == "annee_premier_utilisation":
            m = _re.search(r"(?<!\d)((?:19[89]|20[0-2])\d)(?!\d)", v)
            if m and int(m.group(1)) <= _dt.date.today().year:
                out[f] = m.group(1)
    # الطراز بداية الرقم التسلسلي عادةً: نصحّح به I/O
    fix = str.maketrans("IOQ", "100")
    if out.get("type") and len(out.get("n_chasis", "")) == 17 and \
            out["n_chasis"].startswith(out["type"].translate(fix)):
        out["type"] = out["type"].translate(fix)
    return out


def extract_taxi(text: str, wilaya: int = 32, tsv: str | None = None) -> dict:
    """text: نص الصفحات مفصولة بـ \\f.  tsv (اختياري): مخرجات Tesseract TSV
    للصفحات نفسها مفصولة بـ \\f — تتيح قراءة جدول البطاقة الرمادية بالمواقع."""
    pages = [_clean(p) for p in _re.split(r"\f", text or "") if p.strip()]
    if not pages:
        pages = [""]
    kinds = [_kind(p) for p in pages]

    def pg(kind):
        return [p for p, k in zip(pages, kinds) if kind in k]

    def first(rx, srcs, flags=0, grp=1):
        for s in srcs:
            m = _re.search(rx, s, flags)
            if m:
                return m.group(grp).strip()
        return None

    op, ch, ve, lm = {}, {}, {}, {}

    # ① رخصة الاستغلال المؤقتة ───────────────────────────────────────
    lic = pg("lic") or pages
    for p in lic:
        m = _re.search(r"اللقب\s*[:\-]\s*([^\n]+)", p)
        if m and not op.get("nom_ar"):
            op["nom_ar"] = _ar_words(m.group(1), 3)
        m = _re.search(r"(?:^|\n)[^\n]{0,4}ال[اإ]سم\s*[:\-]\s*([^\n]+)", p)
        if m and not op.get("prenom_ar"):
            op["prenom_ar"] = _ar_words(m.group(1), 3)
        m = _re.search(r"العنوان\s*[:\-]\s*([^\n]+)", p)
        if m and not op.get("adresse_ar"):
            op["adresse_ar"] = _ar_words(m.group(1), 5)
        m = _re.search(r"الصفة\s*[:\-]\s*(\S+)", p)
        if m and not op.get("qualite_oper_physique"):
            q = m.group(1)
            if "مكتر" in q or "مستأجر" in q:
                op["qualite_oper_physique"] = "Locatiare"
            elif _re.search(r"مستفيد|صاحب|مالك|ذو", q):
                op["qualite_oper_physique"] = "Bénéféciaire"
        m = _re.search(r"محيط[^\n:]*:\s*([^\n]+)", p) or \
            _re.search(r"سيارة\s*[اأ]جرة\s*(فردية|ما\s*بين\s*\S+)", p)
        if m and not op.get("nature_activite"):
            v = m.group(1)
            if "فردي" in v:
                op["nature_activite"] = "Individuelle"
            elif "الولاي" in v:
                op["nature_activite"] = "Inter Wilaya"
            elif "البلدي" in v:
                op["nature_activite"] = "Inter Communalle"
        # سطراً سطراً — بعض الأسطر تأتي معكوسة («2245 : رقم الباب»)
        for line in p.split("\n"):
            if _re.search(r"رقم\s*الباب", line) and not op.get("n-porte"):
                m = _re.search(r"رقم\s*الباب\s*[:\-]?\s*(\d{1,5})\b", line) or \
                    _re.search(r"(?<![\d/.])(\d{1,5})\s*[:\-]\s*رقم\s*الباب", line)
                if m:
                    op["n-porte"] = m.group(1)
            if _re.search(r"رقم\s*التسجيل", line) and not ve.get("n_immatriculation"):
                m = _re.search(r"(?<!\d)(\d{4,6})[.\-\s](\d{3})[.\-\s](\d{2})(?!\d)", line)
                if m:
                    ve["n_immatriculation"] = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
            if _re.search(r"انتهاء|غاية", line) and not op.get("date_fin_contrat"):
                d = _dates(line, 2020, 2040)
                if d:
                    op["date_fin_contrat"] = d[0]
        m = _re.search(r"حرر\S*\s*ب\S*\s*في\s*[:.]*\s*([\d/.\-]{8,10})", p)
        if m and not op.get("date_debut_contrat"):
            d = _dates(m.group(1), 2020, 2040)
            if d:
                op["date_debut_contrat"] = d[0]
        # رقم الرخصة «الرقم: 43/2026» ← 2026/43
        m = _re.search(r"الر\s?قم[^\n\d]{0,8}([\d/ \t.|\]\[]{3,20})", p)
        if m and not ch.get("n_delivrance_licence_exploitation"):
            grp = _re.findall(r"\d+", m.group(1))
            year = next((g for g in grp if _re.fullmatch(r"202\d", g)), None)
            if year is None and grp and _re.fullmatch(r"202\d\d{1,4}", grp[0]):
                year, grp = grp[0][:4], [grp[0][:4], grp[0][4:]]
            rest = [g for g in grp if g != year]
            if year and len(rest) == 1 and 2 <= len(rest[0]) <= 4:
                ch["n_delivrance_licence_exploitation"] = f"{year}/{int(rest[0])}"

    # إن غابت رخصة الاستغلال: الاسم من شهادة CNAS / CASNOS
    for p in pg("cnas") + pg("casnos") + pg("id"):
        if not op.get("nom_ar"):
            op["nom_ar"] = _ar_words(first(r"الل?قب\s*[:\-]\s*([^\n]+)", [p]) or "", 3)
        if not op.get("prenom_ar"):
            op["prenom_ar"] = _ar_words(first(r"ال[اإ]سم\s*[:\-]\s*([^\n]+)", [p]) or "", 3)

    nom_op = op.get("nom_ar") or ""
    pre_op = op.get("prenom_ar") or ""

    def is_op_page(p: str) -> bool:
        """هل الصفحة تخصّ المتعامل؟ (لقبه أو اسمه الكامل مذكور فيها)"""
        if not nom_op:
            return True
        words = _re.findall(r"[ء-ي]{2,}", p)
        return any(_similar(w, nom_op) >= 0.75 for w in words) or \
            any(_similar(w, pre_op) >= 0.85 for w in words if len(_ar(w)) >= 4)

    # ② تاريخ الميلاد والأبوان والأسماء اللاتينية — من وثائق المتعامل
    dn_votes: dict[str, int] = {}
    op_docs = [p for p in pg("cg") + pg("cnas") + pg("casnos") + pg("form")]
    for p in op_docs:
        for m in _re.finditer(r"(?:" + _DOB + r"|المولود\S*\s*(?:ب\S*|في)|Naissance)[^\n]{0,60}", p):
            for d in _dates(m.group(0), 1940, 2007):
                dn_votes[d] = dn_votes.get(d, 0) + 2
    for p in pg("id"):
        if is_op_page(p):
            for m in _re.finditer(_DOB + r"[^\n]{0,30}", p):
                for d in _dates(m.group(0), 1940, 2007):
                    dn_votes[d] = dn_votes.get(d, 0) + 1

    # رقم التعريف الوطني: نفضّل ما يوافق سنة الميلاد
    all_nins = []
    for p in pages:
        for n in _nins(p):
            if n not in all_nins:
                all_nins.append(n)
    dn_op = max(dn_votes, key=lambda d: dn_votes[d]) if dn_votes else None
    nin_op = None
    if dn_op:
        nin_op = next((n for n in all_nins if _nin_year(n) == int(dn_op[:4])), None)
        if nin_op is None:
            loose = [n for p in pages for n in _nins_loose(p)
                     if _nin_year(n) == int(dn_op[:4])]
            if len(set(loose)) == 1:
                nin_op = loose[0]
    else:
        # بلا تاريخ ميلاد: نقبل الرقم فقط من بطاقة تعريف تحمل لقب المتعامل
        own = [n for p in pg("id") if is_op_page(p) for n in _nins(p)]
        if len(set(own)) == 1 and 1940 <= (_nin_year(own[0]) or 0) <= 2007:
            nin_op = own[0]
    if nin_op:
        op["nin"] = nin_op
    if dn_op:
        op["date_naissance"] = dn_op

    for p in pg("id"):
        if is_op_page(p) or (nin_op and nin_op in _re.sub(r"[\s.]", "", p)):
            m = _re.search(r"الجنس\s*[:.]?\s*(\S+)", p)
            if m and not op.get("sexe"):
                if _re.search(r"ذكر|زكر|ذكز", m.group(1)):
                    op["sexe"] = "M"
                elif _re.search(r"[اأ]نث", m.group(1)):
                    op["sexe"] = "F"
            m = _re.search(r"Rh\s*[:.]?\s*(AB|A|B|O|0)\s*([+\-#])", p)
            if m and not ch.get("groupage"):
                g = m.group(1).replace("0", "O") + ("-" if m.group(2) == "-" else "+")
                ch["groupage"] = _GROUPAGE.get(g)

    for p in pg("cnas") + pg("casnos") + pg("cg"):
        if not is_op_page(p):
            continue
        if not op.get("nom"):
            op["nom"] = _latin_after(r"Nom\s*", p, 2)
        if not op.get("prenom"):
            op["prenom"] = _latin_after(r"Pr[ée]nom\S*", p, 2)
        if not op.get("prenom_pere"):
            op["prenom_pere"] = _latin_after(r"Fils\s*(?:\(Fille\)\s*)?de", p, 2)
        m = _re.search(r"(?:ابن|أبن|اين|إن)\s*(?:\(ة\))?\s*:\s*([^\n]+)", p)
        m2 = _re.search(r"Et\s*de\s*[:.]?\s*([^\n]+)", p)
        if m and not op.get("prenom_pere_ar"):
            line = m.group(1)
            if m2 and " و " not in line:
                line = line + " و " + m2.group(1)
            pere, nm, pm = _parents(line, nom_op)
            if pere:
                op["prenom_pere_ar"] = pere
            if nm and not op.get("nom_mere_ar"):
                op["nom_mere_ar"], op["prenom_mere_ar"] = nm, pm

    # CG بالفرنسية «M. BENKHIRA Ali»
    if not op.get("nom"):
        for p in pg("cg"):
            m = _re.search(r"\bM[.\s]+([A-Z]{3,})\s+([A-Z][A-Za-z]{2,})", p)
            if m:
                op["nom"], op["prenom"] = _title(m.group(1)), _title(m.group(2))
                break

    # CNAS: رقم الضمان (12 خانة) ورقم عقد الميلاد
    for p in pg("cnas"):
        if not is_op_page(p):
            continue
        m = _re.search(r"(?:تحت\s*رقم|Sous\s*le\s*num\S*)\s*:?\s*([\d/\s]{10,16})", p)
        if m and not op.get("numero_cnas"):
            d = _re.sub(r"\D", "", m.group(1))
            if len(d) == 12 and not _re.search(r"غير\s*منتسب\s*(?:إلى|الى|ل)\s*هي", p[:m.start()]):
                op["numero_cnas"] = d
        m = _re.search(r"(?:عقد\s*الميلاد|N°\s*Acte)\s*:?\s*(\d{2,5})\b", p)
        if m and not op.get("n_acte_naissance"):
            op["n_acte_naissance"] = str(int(m.group(1)))
    for p in pg("casnos"):
        m = _re.search(r"\b(2323\d{7})\b", p)
        if m and not _re.search(r"غير\s*منتسب|NON\s*AFFILI|N.EST\s*PAS", p):
            op["numero_casnos"] = m.group(1)
    for p in pg("cnas") + pg("casnos"):
        m = _re.search(r"Fait\s*le\s*:?\s*(\d{2}/\d{2}/\d{4})", p)
        if m:
            d = _dates(m.group(1), 2023, 2030)
            if d and (not ch.get("date_mise_a_jour_securite_sociale")
                      or d[0] > ch["date_mise_a_jour_securite_sociale"]):
                ch["date_mise_a_jour_securite_sociale"] = d[0]

    # ③ البطاقة الرمادية ─────────────────────────────────────────────
    # أولاً: القراءة بالمواقع — القيمة هي ما فوق العنوان المطبوع في الخانة نفسها
    tsv_pages = tsv.split("\f") if tsv else []
    for i, k in enumerate(kinds):
        if "cg" in k and i < len(tsv_pages) and tsv_pages[i].strip():
            for f, v in _cg_cells(tsv_pages[i]).items():
                ve.setdefault(f, v)

    for p in pg("cg"):
        up = p.upper()
        # الصنف: نبحث في سطر القيم فوق «GENRE MARQUE» فقط — كلمة «KIA» وسط
        # ضجيج الصفحة لا تعني أن المركبة كيا
        ls = up.split("\n")
        idx = [i for i, l in enumerate(ls) if "MARQUE" in l or "الصنف" in l]
        zone = "\n".join("\n".join(ls[max(0, i - 2): i + 1]) for i in idx)
        if not ve.get("marque") and zone:
            ve["marque"] = _brand(zone, short_ok=True)
        if not ve.get("marque"):
            ve["marque"] = _brand(up, short_ok=False)
        if not ve.get("marque"):
            ve.pop("marque", None)
        if not ve.get("marque"):
            for ar, mk in _MARQUES_AR.items():
                if ar in p:
                    ve["marque"] = mk
                    break
        # الطراز + الرقم التسلسلي (VIN) على السطر نفسه بعد الصنف
        if ve.get("marque"):
            m = _re.search(ve["marque"][:4] + r"\S*\s*[|]?\s*([A-Z0-9]{5,12})\s*[|]?\s*([A-Z0-9]{8,17})?", up)
            if m:
                typ, vin = m.group(1), m.group(2)
                if vin and len(vin) == 17:
                    fix = str.maketrans("IOQ", "100")      # لا I ولا O ولا Q في VIN
                    vin = vin.translate(fix)
                    if typ and vin.startswith(typ.translate(fix)):
                        typ = typ.translate(fix)
                if typ and _re.search(r"[A-Z]", typ) and _re.search(r"\d", typ):
                    ve.setdefault("type", typ)
                if vin and len(vin) in (8, 17) and _re.search(r"\d", vin):
                    ve.setdefault("n_chasis", vin)
        # الهيكل · الطاقة · القوة · المقاعد
        m = _re.search(r"\b(BREAK|BERLINE|BRIAK|REAK|BERLIN|C\.?I)\b[,.]?\s*(EG|ES|GO|GPL|ESS\S*)?"
                       r"\s*(\d{1,2})?\s*(\d)?", up)
        c = _carrosserie(m.group(1)) if m else \
            ("BERLINE" if "برلين" in p else "C,I" if _re.search(r"ق\s?\.\s?د", p) else None)
        if c:
            ve.setdefault("carrossere", c)
        if m and m.group(2):
            e = m.group(2)
            ve.setdefault("energie", "Diesel" if e == "GO" else
                          "GPL" if e == "GPL" else "Essence")   # EG = Essence (كما في ملفاتك)
            if m.group(3) and 3 <= int(m.group(3)) <= 15:
                ve.setdefault("puissance", f"{int(m.group(3))} CHEVEAUX")
            if m.group(4) and 2 <= int(m.group(4)) <= 9:
                ve.setdefault("n_siege", m.group(4))
        if not ve.get("energie"):
            if _re.search(r"بنز\S*\s*\+\s*غاز|\+\s*غاز", p):
                ve["energie"] = "GPL"
            elif _re.search(r"مازوت|\bGO\b|GASOIL", p):
                ve["energie"] = "Diesel"
        if _re.search(r"سياح|س\s?[.,]?\s?خ\b|\bVP\b", p):
            ve.setdefault("genre", "س,خ")
        # سنة أول استعمال: سنة 19xx/20xx بجانب رقمي التسجيل
        for m in _re.finditer(r"(?<!\d)((?:19[89]|20[0-2])\d)(?!\d)", p):
            ctx = p[max(0, m.start() - 40): m.end() + 40]
            if _re.search(r"\d{4,6}-\d{2,3}-?\d{1,2}", ctx) and int(m.group(1)) <= _dt.date.today().year:
                ve.setdefault("annee_premier_utilisation", m.group(1))
                break
        if not ve.get("n_immatriculation"):
            for m in _re.finditer(r"(?<!\d)(\d{5,6})-(\d{3})-(32)(?!\d)", p):
                ve["n_immatriculation"] = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
                break
        m = _re.search(r"(?:ابن|أبن)\s*:\s*([^\n]+)", p)
        if m and is_op_page(p) and not op.get("prenom_pere_ar"):
            pere, nm, pm = _parents(m.group(1), nom_op)
            if pere:
                op["prenom_pere_ar"] = pere
            if nm:
                op["nom_mere_ar"], op["prenom_mere_ar"] = nm, pm

    # ④ شهادة التكوين ───────────────────────────────────────────────
    for p in pg("form"):
        head = "\n".join(l for l in p.split("\n")[:12]
                         if not _re.search(r"مرسوم|المرصوم|المرسوع|قرار|مقتضى", l))
        m = _re.search(r"رق\S*\s*\D{0,3}(\d{2,5})\b[^\n]*?\b(202\d)\b", head)
        if m and not ch.get("n_certificat_formation"):
            ch["n_certificat_formation"] = f"{m.group(2)}/{int(m.group(1))}"
        for m in _re.finditer(r"(?:البي\S*|في)\s*[:;,.]?\s*\S{0,6}\s*[;:,]?\s*"
                              r"(\d{4}[/\-]\d{2}[/\-]\d{2})", head):
            d = _dates(m.group(1), 2015, 2035)
            if d:
                ch.setdefault("date_certificat_formation", d[0])
                break

    # ⑤ رخصة السياقة (نادراً ما تُقرأ — صورة بطاقة) ───────────────────
    for p in pg("permis"):
        m = _re.search(r"\b(L\s?\d{8})\b|\b(\d{2}/\d{2}/\d{5})\b", p)
        if m and not ch.get("n_permis"):
            ch["n_permis"] = (m.group(1) or m.group(2)).replace(" ", "")

    # ⑥ قرار منح الرخصة لذي الحق (المجاهد/الأرملة/الابنة) ──────────────
    _STOP = {"ابنة", "إبنة", "ابن", "إبن", "بنت", "ابت", "ابئة", "المولود", "المولودة",
             "بصفتها", "بصفته", "واليا", "والي", "أرملة", "ارملة", "مكان"}
    _name_rx = (r"للسيد(?:\(?ة\)?|زة\)|ر\(ة\)|ه)?\s*[:؛)]*\s*"
                r"([ء-ي]{2,})\s+([ء-ي]{2,})(?:\s+([ء-ي]{2,}))?")
    for p in pg("dec"):
        all_l = p.split("\n")
        lines = [(i, l) for i, l in enumerate(all_l)
                 if "للسيد" in l and not _re.search(r"تعيين|زعيين|واليا", l)]
        # الأولوية لسطر «المادة الأولى: تمنح رخصة ... للسيد(ة)» ثم متن القرار ثم العنوان
        lines.sort(key=lambda il: (0 if _re.search(r"ادة\s*الا?و?لى|تمنح|تمتح|تملع", il[1])
                                   else 1 if il[0] >= 8 else 2, il[0]))
        lines = [l for _, l in lines]
        for l in lines:
            m = _re.search(_name_rx, l)
            if not m:
                continue
            nom_b, w2, w3 = m.group(1), m.group(2), m.group(3)
            if nom_b in _STOP or w2 in _STOP:
                continue
            pre_b = w2 + " " + w3 if (w2 in ("عبد", "أبو", "ابو") and w3) else w2
            lm.setdefault("nom_ar", nom_b)
            lm.setdefault("prenom_ar", pre_b)
            break
        m = _re.search(r"المولود\S*\s*بتاريخ\s*[:]?\s*([\d/.\-]{8,10})", p)
        if m:
            d = _dates(m.group(1), 1900, 2007)
            if d:
                lm.setdefault("date_naissance", d[0])
        m = _re.search(r"مكان\s*الاست\S*\s*(?:ب)?\S*\s+([^\n؛;]+)", p)
        if m:
            c = _commune(m.group(1))
            if c:
                lm.setdefault("concentration_exploitation_licence_modjahid", c)
        for l in p.split("\n")[:12]:
            if _re.search(r"مقتضى|بناء|بتاء|على\s*القرار", l):
                continue           # إحالة إلى نص قانوني آخر — ليس رقم هذا القرار
            m = _re.search(r"قرار\s*رقم\s*:?\s*(\d{2,5})(?![\d/])(?!\s+\d)", l)
            if m:
                lm.setdefault("n_decision_octroi", m.group(1))
                d = _dates(l[m.end():], 1962, 2030)
                if d:
                    lm.setdefault("date_decision_octroi", d[0])
                break
        if _re.search(r"تمنح|منح", p):
            lm.setdefault("type_licence", "Délivrée")

    # بطاقة ذي الحق: الرقم الوطني الذي يوافق اسمه أو سنة ميلاده
    if lm:
        by = int(lm["date_naissance"][:4]) if lm.get("date_naissance") else None
        for p in pg("id"):
            if is_op_page(p) and not (lm.get("nom_ar") and lm["nom_ar"] in p):
                continue
            ns = [n for n in _nins(p) if n != nin_op]
            for n in ns:
                if by is None or _nin_year(n) == by:
                    lm.setdefault("nin_modjahid", n)
            if lm.get("nin_modjahid") and not lm.get("date_naissance"):
                for m in _re.finditer(_DOB + r"[^\n]{0,30}", p):
                    d = _dates(m.group(0), 1900, 2007)
                    if d and int(d[0][:4]) == _nin_year(lm["nin_modjahid"]):
                        lm["date_naissance"] = d[0]
                        break
        if not lm.get("nin_modjahid"):
            acte_op = op.get("n_acte_naissance")

            def plausible(n):
                if n == nin_op:
                    return False
                if acte_op and str(int(n[9:14])) == acte_op:
                    return False   # رقم المتعامل نفسه مقروءاً بخطأ
                if by is not None:
                    return _nin_year(n) == by
                return n[1] == "1" or (_nin_year(n) or 3000) <= 1975
            rest = [n for n in all_nins if plausible(n)]
            if len(rest) == 1:
                lm["nin_modjahid"] = rest[0]

    # ⑦ التجميع والنسخ بين الأقسام ───────────────────────────────────
    porte = op.get("n-porte")
    nature = op.get("nature_activite")
    op.update({"id_wilaya": str(wilaya), "activite": "Transports par Taxi"})
    if op.get("date_debut_contrat") and op.get("date_fin_contrat") and \
            op["date_debut_contrat"] >= op["date_fin_contrat"]:
        op.pop("date_debut_contrat")

    ch.update({
        "id_wilaya": str(wilaya), "activite": "Transports par Taxi",
        "nature_activite": nature, "qualite_chauffeur": "Principal",
        "nin": op.get("nin"), "nom_ar": op.get("nom_ar"), "prenom_ar": op.get("prenom_ar"),
        "nom": op.get("nom"), "prenom": op.get("prenom"),
        "date_naissance": op.get("date_naissance"), "sexe": op.get("sexe"),
        "n_acte_naissance": op.get("n_acte_naissance"),
        "numero_cnas": op.get("numero_cnas"), "numero_casnos": op.get("numero_casnos"),
        "adresse_ar": op.get("adresse_ar"),
    })
    if ch.get("n_delivrance_licence_exploitation") and op.get("date_debut_contrat"):
        ch.setdefault("date_delivrance_licence_exploitation", op["date_debut_contrat"])

    ve.update({
        "id_wilaya": str(wilaya), "activite": "Transports par Taxi",
        "nature_activite": nature, "n_porte": porte,
        "adresse_ar": op.get("adresse_ar"),
    })
    if ch.get("n_delivrance_licence_exploitation"):
        ve.setdefault("n_certificat_cerculation", ch["n_delivrance_licence_exploitation"])
    if op.get("date_fin_contrat"):
        ve.setdefault("date_fin_validite_certificat", op["date_fin_contrat"])

    if lm:
        lm.update({"id_wilaya": str(wilaya), "n_porte": porte})

    def tidy(d):
        return {k: v for k, v in d.items() if v not in (None, "", [])}

    return {
        "operateur": tidy(op),
        "chauffeur": tidy(ch),
        "vehicule": tidy(ve),
        "licence_modjahid": tidy(lm),
    }


# ──────────────────────────────────────────────────────────────────
#  معالجة ملف PDF واحد
# ──────────────────────────────────────────────────────────────────

def process_pdf(pdf_path: Path, tess_exe: str, tessdata: str | None) -> dict:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz

    doc = fitz.open(str(pdf_path))
    lang = pick_lang(tessdata)
    print(f"  📄 {len(doc)} صفحة [{lang}]", end="", flush=True)

    pages_text: list[str] = []
    pages_tsv: list[str] = []
    for i, page in enumerate(doc):
        # نجرب النص المدمج أولاً
        embedded = page.get_text().strip()
        if embedded:
            pages_text.append(embedded)
            pages_tsv.append("")
            continue
        # صورة ممسوحة → Tesseract
        pix = page.get_pixmap(dpi=DPI)
        img_bytes = pix.tobytes("jpeg")
        txt, tsv = ocr_image(img_bytes, tess_exe, tessdata, lang, with_tsv=True)
        pages_text.append(txt.replace("\f", "\n"))
        pages_tsv.append(tsv)
        print(".", end="", flush=True)
    doc.close()

    print()
    # الصفحات مفصولة بـ \f: كل وثيقة تُحلَّل حسب نوعها
    return extract_taxi("\f".join(pages_text), WILAYA, tsv="\f".join(pages_tsv))

# ──────────────────────────────────────────────────────────────────
#  كتابة في Excel
# ──────────────────────────────────────────────────────────────────

def write_excel(data: dict, excel_path: Path):
    try:
        import openpyxl
    except ImportError:
        print("  ⚠️  openpyxl غير مثبّت — لن يُكتب في Excel")
        return

    if not excel_path.exists():
        print(f"  ⚠️  ملف Excel غير موجود: {excel_path}")
        return

    wb = openpyxl.load_workbook(str(excel_path))
    changed = False

    sheet_map = {
        "operateur": "Operateur",
        "chauffeur":  "Chauffeur",
        "vehicule":   "Véhicule",
        "licence_modjahid": "Licence Modjahid",
    }

    for grp, sheet_name in sheet_map.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        row_data = data.get(grp, {})
        if not any(v for v in row_data.values() if v is not None):
            continue

        # أول صف فارغ من الصف 3
        first_empty = ws.max_row + 1
        for r in range(3, ws.max_row + 2):
            if all(ws.cell(r, c).value is None for c in range(1, min(ws.max_column + 1, 10))):
                first_empty = r
                break

        # ترسيم الأعمدة
        headers = {ws.cell(1, c).value or ws.cell(2, c).value: c
                   for c in range(1, ws.max_column + 1)}
        for key, val in row_data.items():
            if val is None:
                continue
            col = headers.get(key)
            if col:
                ws.cell(first_empty, col, val)
                changed = True

    if changed:
        # نسخة احتياطية
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        bk = excel_path.parent / f"{excel_path.stem}_bk_{ts}.xlsx"
        import shutil
        shutil.copy2(str(excel_path), str(bk))
        wb.save(str(excel_path))
        print(f"  ✅ كُتب في Excel (نسخة احتياطية: {bk.name})")
    else:
        print("  ⚠️  لم يُكتب شيء (لم تُعثر على الأعمدة المطابقة)")

# ──────────────────────────────────────────────────────────────────
#  النقطة الرئيسية
# ──────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="استخراج بيانات سيارات الأجرة محلياً")
    parser.add_argument("files", nargs="*", help="ملفات PDF (فارغ = كل المجلد)")
    parser.add_argument("--setup", action="store_true", help="تثبيت اللغة العربية")
    parser.add_argument("--excel", default=str(EXCEL_PATH), help="مسار ملف Excel")
    parser.add_argument("--no-excel", action="store_true", help="لا تكتب في Excel")
    args = parser.parse_args()

    if args.setup:
        setup_arabic()
        return

    tess_exe = find_tesseract()
    tessdata = find_tessdata()

    print(f"🔧 Tesseract: {tess_exe}")
    print(f"📂 tessdata:  {tessdata or '❌ غير موجود — شغّل: python offline_ocr.py --setup'}")
    if not tessdata:
        sys.exit(1)
    lang = pick_lang(tessdata)
    print(f"🔤 اللغات:    {lang}"
          + ("" if "fra" in lang else "   ← شغّل --setup لإضافة الفرنسية (أسماء لاتينية، صنف وطراز المركبة)"))

    # تحديد الملفات
    if args.files:
        pdfs = [Path(p) for p in args.files]
    else:
        pdfs = sorted(SCRIPT_DIR.glob("*.pdf"))

    if not pdfs:
        print("لم يُعثر على ملفات PDF في المجلد.")
        sys.exit(0)

    print(f"\n📋 {len(pdfs)} ملف للمعالجة\n{'─' * 50}")

    excel_path = Path(args.excel)
    results = []

    for i, pdf in enumerate(pdfs, 1):
        print(f"[{i}/{len(pdfs)}] {pdf.name}")
        try:
            data = process_pdf(pdf, tess_exe, tessdata)
            # حفظ JSON جانبي
            json_out = pdf.with_suffix(".json")
            json_out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"  💾 {json_out.name}")

            # عرض ملخص
            op = data.get("operateur", {})
            print(f"  👤 {op.get('nom_ar') or '—'}  {op.get('prenom_ar') or '—'}"
                  f"  | باب: {op.get('n-porte') or '—'}"
                  f"  | انتهاء: {op.get('date_fin_contrat') or '—'}")
            ve = data.get("vehicule", {})
            print(f"  🚙 {ve.get('n_immatriculation') or '—'}  {ve.get('marque') or '—'}"
                  f"  {ve.get('type') or ''}  | سنة: {ve.get('annee_premier_utilisation') or '—'}"
                  f"  | طاقة: {ve.get('energie') or '—'}")
            filled = sum(len(v) for v in data.values() if isinstance(v, dict))
            print(f"  📊 {filled} حقلاً مستخرجاً")

            if not args.no_excel:
                write_excel(data, excel_path)

            results.append({"file": pdf.name, "status": "ok", "data": data})
        except Exception as e:
            print(f"  ❌ خطأ: {e}")
            results.append({"file": pdf.name, "status": "error", "error": str(e)})

    # ملخص نهائي
    ok    = sum(1 for r in results if r["status"] == "ok")
    fails = len(results) - ok
    print(f"\n{'═' * 50}")
    print(f"✅ نجح: {ok}  |  ❌ فشل: {fails}  |  المجموع: {len(results)}")


if __name__ == "__main__":
    main()
